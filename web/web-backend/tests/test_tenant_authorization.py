import asyncio
import os
import sys
from datetime import date, datetime
from io import BytesIO
from types import SimpleNamespace

import openpyxl
import pytest
from fastapi import HTTPException, UploadFile
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import models
from database import Base
from database import get_db
from main import app, assert_question_access, validate_account_request_scope, create_department, create_subject_manually, update_user_department, get_admin_user_overview, DepartmentCreateRequest, UserDepartmentUpdateRequest, SubjectCreateRequest, get_subjects, get_faculty_program_subjects, get_questions, get_super_admin_overview, get_super_admin_ai_usage, get_admin_campus_overview, get_admin_ai_usage, get_question_sets, export_question_bank_tos, upload_files
from main import get_question_versions
from routers.questions import _assert_upload_access, _resolve_department_leadership, export_institutional_tos, upload_and_analyze_syllabus
from routers.activity import get_activity_logs
from security import assert_campus_access, assert_user_subject_campus_access, get_current_user, require_admin, require_campus_admin, require_super_admin, visible_campus_id


@pytest.fixture
def db_session():
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    try:
        yield session
    finally:
        session.close()
        engine.dispose()


def test_assessment_analytics_requires_auth_and_scopes_tos_and_reuse_candidates():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    session_factory = sessionmaker(bind=engine)
    session = session_factory()
    campus = models.Campus(name="Analytics Campus", code="ANALYTICS")
    session.add(campus)
    session.flush()
    faculty = models.User(
        email="analytics_faculty@example.com",
        password="hashed",
        role="faculty",
        campus_id=campus.id,
    )
    other_faculty = models.User(
        email="analytics_other@example.com",
        password="hashed",
        role="faculty",
        campus_id=campus.id,
    )
    session.add_all([faculty, other_faculty])
    session.flush()

    def add_tos(owner, subject_name, tos_data):
        subject = models.Subject(name=subject_name, code=subject_name[:8], user_id=owner.id)
        session.add(subject)
        session.flush()
        upload = models.UploadedFile(user_id=owner.id, subject_id=subject.id)
        session.add(upload)
        session.flush()
        tos = models.TableOfSpecification(
            upload_id=upload.id,
            tos_data=tos_data,
            total_items=1,
            exam_type="Quiz",
        )
        session.add(tos)
        session.flush()
        return tos

    target = add_tos(
        faculty,
        "Target Course",
        [{"topic_name": "Database Design", "bloom_counts": {"Apply": 1}}],
    )
    accessible_source = add_tos(faculty, "Other Course", [])
    hidden_source = add_tos(other_faculty, "Hidden Course", [])
    session.add_all([
        models.GeneratedQuestion(
            tos_id=accessible_source.id,
            user_id=faculty.id,
            topic_name="Database Design",
            bloom_level="Apply",
            question_type="MCQ",
            question="Which design best avoids duplicate data?",
        ),
        models.GeneratedQuestion(
            tos_id=hidden_source.id,
            user_id=other_faculty.id,
            topic_name="Database Design",
            bloom_level="Apply",
            question_type="MCQ",
            question="Hidden question from another faculty.",
        ),
    ])
    session.commit()

    def override_db():
        yield session

    previous_overrides = app.dependency_overrides.copy()
    app.dependency_overrides[get_db] = override_db
    try:
        with TestClient(app) as client:
            unauthenticated = client.get("/api/analytics/tos-list")
            assert unauthenticated.status_code == 401

            app.dependency_overrides[get_current_user] = lambda: faculty
            records = client.get("/api/analytics/tos-list")
            assert records.status_code == 200
            assert {item["tos_id"] for item in records.json()} == {
                target.id,
                accessible_source.id,
            }

            forecast = client.get(f"/api/analytics/tos/{target.id}/forecast")
            assert forecast.status_code == 200
            forecast_data = forecast.json()["predictive"]
            assert forecast_data["overall_completion_pct"] == 0
            assert forecast_data["overall_target_total"] == 1
            assert forecast_data["at_risk_bloom_levels"] == []

            suggestions = client.get(
                f"/api/analytics/tos/{target.id}/reuse-suggestions"
            )
            assert suggestions.status_code == 200
            matches = suggestions.json()["suggestions"][0]["matches"]
            assert [match["question"] for match in matches] == [
                "Which design best avoids duplicate data?"
            ]

            hidden_forecast = client.get(
                f"/api/analytics/tos/{hidden_source.id}/forecast"
            )
            assert hidden_forecast.status_code == 404
    finally:
        app.dependency_overrides.clear()
        app.dependency_overrides.update(previous_overrides)
        session.close()
        engine.dispose()


def test_super_admin_can_access_any_campus():
    user = SimpleNamespace(role="super_admin", campus_id=None)

    assert assert_campus_access(user, 12) is None
    assert visible_campus_id(user) is None
    assert require_admin(user) is user
    assert require_super_admin(user) is user


def test_campus_admin_is_limited_to_assigned_campus():
    user = SimpleNamespace(role="campus_admin", campus_id=4)

    assert assert_campus_access(user, 4) is None
    assert visible_campus_id(user) == 4
    assert require_admin(user) is user
    assert require_campus_admin(user) is user

    with pytest.raises(HTTPException) as error:
        assert_campus_access(user, 5)

    assert error.value.status_code == 403


def test_legacy_admin_is_limited_to_assigned_campus():
    user = SimpleNamespace(role="admin", campus_id=7)

    assert assert_campus_access(user, 7) is None
    assert visible_campus_id(user) == 7
    assert require_admin(user) is user
    assert require_campus_admin(user) is user

    with pytest.raises(HTTPException) as error:
        assert_campus_access(user, 8)

    assert error.value.status_code == 403


def test_legacy_admin_without_campus_assignment_is_rejected():
    user = SimpleNamespace(role="admin", campus_id=None)

    with pytest.raises(HTTPException) as error:
        require_admin(user)

    assert error.value.status_code == 403


def test_campus_admin_without_assignment_is_rejected():
    user = SimpleNamespace(role="campus_admin", campus_id=None)

    with pytest.raises(HTTPException) as error:
        visible_campus_id(user)

    assert error.value.status_code == 403
    with pytest.raises(HTTPException) as admin_error:
        require_admin(user)

    assert admin_error.value.status_code == 403


def test_faculty_cannot_use_admin_campus_scope():
    user = SimpleNamespace(role="faculty", campus_id=4)

    with pytest.raises(HTTPException) as admin_error:
        require_admin(user)
    with pytest.raises(HTTPException) as campus_error:
        assert_campus_access(user, 4)
    with pytest.raises(HTTPException) as campus_role_error:
        require_campus_admin(user)

    assert admin_error.value.status_code == 403
    assert campus_error.value.status_code == 403
    assert campus_role_error.value.status_code == 403


def test_faculty_removal_hides_shared_subject_only_for_that_user(db_session):
    from main import delete_subject, restore_subject

    creator = models.User(email="subject-creator@example.com", password="test", role="faculty")
    remover = models.User(email="subject-remover@example.com", password="test", role="faculty")
    other_faculty = models.User(email="subject-other@example.com", password="test", role="faculty")
    db_session.add_all([creator, remover, other_faculty])
    db_session.flush()

    subject = models.Subject(name="Shared Subject Removal Test", user_id=creator.id)
    db_session.add(subject)
    db_session.flush()
    for faculty in (remover, other_faculty):
        db_session.add(models.UploadedFile(
            user_id=faculty.id,
            subject_id=subject.id,
            module_filename="module.pdf",
            syllabus_filename="cis.pdf",
            module_text="module text",
            syllabus_text="CIS text",
        ))
        db_session.add(models.GeneratedQuestion(
            subject_id=subject.id,
            user_id=faculty.id,
            question=f"Question for {faculty.email}",
        ))
    db_session.commit()

    result = delete_subject(subject.id, user_id=remover.id, db=db_session, current_user=remover)

    assert result["message"] == "Subject removed from your Question Bank."
    assert db_session.query(models.Subject).filter(models.Subject.id == subject.id).one().archived is False
    assert subject.id not in [item["id"] for item in get_subjects(user_id=remover.id, db=db_session, current_user=remover)]
    assert subject.id in [item["id"] for item in get_subjects(user_id=other_faculty.id, db=db_session, current_user=other_faculty)]
    assert get_questions(subject_id=subject.id, user_id=remover.id, db=db_session, current_user=remover) == []
    assert len(get_questions(subject_id=subject.id, user_id=other_faculty.id, db=db_session, current_user=other_faculty)) == 1

    restore_subject(subject.id, user_id=remover.id, db=db_session, current_user=remover)

    assert subject.id in [item["id"] for item in get_subjects(user_id=remover.id, db=db_session, current_user=remover)]
    assert db_session.query(models.Subject).filter(models.Subject.id == subject.id).one().archived is False


def test_faculty_profile_updates_cannot_change_program_directly(db_session):
    from main import UserProfileUpdateRequest, get_user_profile, update_user_profile

    campus = models.Campus(name="Profile Settings Campus", code="PROFILE-SETTINGS")
    db_session.add(campus)
    db_session.flush()
    department = models.Department(name="Profile Computing", campus_id=campus.id)
    other_department = models.Department(name="Profile Engineering", campus_id=campus.id)
    db_session.add_all([department, other_department])
    db_session.flush()

    current_program = models.Program(name="Current Program", department_id=department.id)
    next_program = models.Program(name="Next Program", department_id=department.id)
    other_program = models.Program(name="Other Department Program", department_id=other_department.id)
    db_session.add_all([current_program, next_program, other_program])
    db_session.flush()

    faculty = models.User(
        email="profile-faculty@example.com",
        password="test",
        role="faculty",
        name="Current Faculty",
        department="Stale Department Label",
        campus_id=campus.id,
        program_id=current_program.id,
    )
    db_session.add(faculty)
    db_session.commit()

    profile = get_user_profile(db=db_session, current_user=faculty)

    assert profile["department"] == department.name
    assert profile["program_id"] == current_program.id
    assert profile["program_name"] == current_program.name
    assert profile["role"] == "faculty"
    assert {program["id"] for program in profile["programs"]} == {current_program.id, next_program.id}

    updated = update_user_profile(
        UserProfileUpdateRequest(full_name="Updated Faculty"),
        db=db_session,
        current_user=faculty,
    )

    assert updated["full_name"] == "Updated Faculty"
    assert updated["program_id"] == current_program.id
    assert faculty.department == "Stale Department Label"
    with pytest.raises(HTTPException) as error:
        update_user_profile(
            UserProfileUpdateRequest(full_name="Updated Faculty", program_id=other_program.id),
            db=db_session,
            current_user=faculty,
        )
    assert error.value.status_code == 403
    assert faculty.program_id == current_program.id


def test_faculty_program_change_request_keeps_exact_target_until_admin_approval(db_session):
    from main import (
        UserChangeRequestPayload,
        UserChangeReviewPayload,
        create_user_change_request,
        list_user_change_requests,
        review_user_change_request,
    )

    campus = models.Campus(name="Change Request Campus", code="CHANGE-REQUEST")
    db_session.add(campus)
    db_session.flush()
    current_department = models.Department(name="Current Department", campus_id=campus.id)
    target_department = models.Department(name="Target Department", campus_id=campus.id)
    db_session.add_all([current_department, target_department])
    db_session.flush()
    current_program = models.Program(name="Shared Program Name", department_id=current_department.id)
    target_program = models.Program(name="Shared Program Name", department_id=target_department.id)
    db_session.add_all([current_program, target_program])
    db_session.flush()
    faculty = models.User(
        email="change_request_faculty@example.com",
        password="not-used",
        role="faculty",
        department=current_department.name,
        campus_id=campus.id,
        program_id=current_program.id,
    )
    admin = models.User(email="change_request_admin@example.com", password="not-used", role="super_admin")
    db_session.add_all([faculty, admin])
    db_session.commit()

    result = create_user_change_request(
        UserChangeRequestPayload(
            user_id=faculty.id,
            request_type="program",
            requested_value=str(target_program.id),
        ),
        db=db_session,
        current_user=faculty,
    )
    request = db_session.query(models.UserChangeRequest).filter(
        models.UserChangeRequest.id == result["id"]
    ).one()
    assert faculty.program_id == current_program.id
    assert request.requested_value == str(target_program.id)

    listed = list_user_change_requests(status="pending", db=db_session, _admin=admin)
    assert listed[0]["requested_value"] == "Shared Program Name / Target Department"

    review_user_change_request(
        request.id,
        UserChangeReviewPayload(action="approve"),
        db=db_session,
        admin=admin,
    )
    assert faculty.program_id == target_program.id
    assert faculty.department == target_department.name


def test_question_access_resolves_campus_through_academic_hierarchy(db_session):
    lipa = models.Campus(name="Lipa", code="LIPA")
    alangilan = models.Campus(name="Alangilan", code="ALANGILAN")
    db_session.add_all([lipa, alangilan])
    db_session.flush()
    department = models.Department(name="Computing", campus_id=lipa.id)
    db_session.add(department)
    db_session.flush()
    program = models.Program(name="Information Technology", department_id=department.id)
    db_session.add(program)
    db_session.flush()
    subject = models.Subject(name="Programming", program_id=program.id, department_id=department.id)
    db_session.add(subject)
    db_session.flush()
    question = models.GeneratedQuestion(subject_id=subject.id, question="What is a program?")
    db_session.add(question)
    db_session.commit()

    lipa_admin = SimpleNamespace(role="campus_admin", campus_id=lipa.id)
    alangilan_admin = SimpleNamespace(role="campus_admin", campus_id=alangilan.id)

    assert assert_question_access(db_session, lipa_admin, question) is None
    with pytest.raises(HTTPException) as error:
        assert_question_access(db_session, alangilan_admin, question)

    assert error.value.status_code == 403


def test_faculty_and_campus_admin_cannot_read_foreign_campus_questions(db_session):
    lipa = models.Campus(name="Lipa Access", code="LIPA-ACCESS")
    main = models.Campus(name="Main Access", code="MAIN-ACCESS")
    db_session.add_all([lipa, main])
    db_session.flush()
    lipa_department = models.Department(name="Lipa Department", campus_id=lipa.id)
    main_department = models.Department(name="Main Department", campus_id=main.id)
    db_session.add_all([lipa_department, main_department])
    db_session.flush()
    lipa_program = models.Program(name="Lipa Program", department_id=lipa_department.id)
    main_program = models.Program(name="Main Program", department_id=main_department.id)
    db_session.add_all([lipa_program, main_program])
    db_session.flush()
    lipa_faculty = models.User(email="lipa_faculty@example.com", password="secret", role="faculty", campus_id=lipa.id, program_id=lipa_program.id)
    main_faculty = models.User(email="main_faculty@example.com", password="secret", role="faculty", campus_id=main.id, program_id=main_program.id)
    lipa_admin = models.User(email="lipa_admin@example.com", password="secret", role="campus_admin", campus_id=lipa.id)
    db_session.add_all([lipa_faculty, main_faculty, lipa_admin])
    db_session.flush()
    shared_subject = models.Subject(name="Shared Access Subject", department_id=lipa_department.id, user_id=lipa_faculty.id)
    db_session.add(shared_subject)
    db_session.flush()
    lipa_question = models.GeneratedQuestion(subject_id=shared_subject.id, user_id=lipa_faculty.id, question="Lipa question")
    main_question = models.GeneratedQuestion(subject_id=shared_subject.id, user_id=main_faculty.id, question="Main question")
    db_session.add_all([lipa_question, main_question])
    foreign_subject = models.Subject(name="Main Campus Upload Subject", department_id=main_department.id)
    db_session.add(foreign_subject)
    db_session.flush()
    foreign_subject_upload = models.UploadedFile(
        user_id=lipa_faculty.id,
        subject_id=foreign_subject.id,
        module_filename="module.pdf",
        syllabus_filename="syllabus.pdf",
        module_text="module",
        syllabus_text="syllabus",
    )
    db_session.add(foreign_subject_upload)
    question_set = models.QuestionSet(name="Shared subject set", subject_id=shared_subject.id)
    db_session.add(question_set)
    db_session.flush()
    db_session.add_all([
        models.QuestionSetItem(question_set_id=question_set.id, question_id=lipa_question.id, position=0),
        models.QuestionSetItem(question_set_id=question_set.id, question_id=main_question.id, position=1),
    ])
    db_session.commit()

    faculty_questions = get_questions(db=db_session, current_user=lipa_faculty)
    admin_questions = get_questions(db=db_session, current_user=lipa_admin)
    faculty_sets = get_question_sets(db=db_session, current_user=lipa_faculty)
    admin_sets = get_question_sets(db=db_session, current_user=lipa_admin)

    assert {item["id"] for item in faculty_questions} == {lipa_question.id}
    assert {item["id"] for item in admin_questions} == {lipa_question.id}
    assert faculty_sets[0]["question_ids"] == [lipa_question.id]
    assert faculty_sets[0]["question_count"] == 1
    assert admin_sets[0]["question_ids"] == [lipa_question.id]
    assert_user_subject_campus_access(db_session, lipa_faculty, shared_subject)
    with pytest.raises(HTTPException) as subject_error:
        assert_user_subject_campus_access(db_session, main_faculty, shared_subject)
    assert subject_error.value.status_code == 403
    with pytest.raises(HTTPException) as error:
        assert_question_access(db_session, lipa_faculty, main_question)
    assert error.value.status_code == 404
    with pytest.raises(HTTPException) as upload_error:
        _assert_upload_access(str(foreign_subject_upload.id), lipa_faculty, db_session)
    assert upload_error.value.status_code == 403


def test_account_request_must_match_selected_campus_department_and_program(db_session):
    lipa = models.Campus(name="Lipa Request", code="LIPA-REQ")
    alangilan = models.Campus(name="Alangilan Request", code="ALA-REQ")
    db_session.add_all([lipa, alangilan])
    db_session.flush()
    department = models.Department(name="Computing Request", campus_id=lipa.id)
    db_session.add(department)
    db_session.flush()
    program = models.Program(name="IT Request", department_id=department.id)
    db_session.add(program)
    db_session.commit()

    assert validate_account_request_scope(
        db_session, lipa.id, department.name, program.id
    ) is None
    with pytest.raises(HTTPException) as error:
        validate_account_request_scope(
            db_session, alangilan.id, department.name, program.id
        )

    assert error.value.status_code == 422


def test_campus_admin_can_list_subjects_by_program_membership(db_session):
    campus = models.Campus(name="Main Campus", code="MAIN")
    db_session.add(campus)
    db_session.flush()
    department = models.Department(name="College of Teaching Education", campus_id=campus.id)
    db_session.add(department)
    db_session.flush()
    program = models.Program(name="Bachelor of Secondary Education", department_id=department.id)
    other_program = models.Program(name="Bachelor of Elementary Education", department_id=department.id)
    db_session.add_all([program, other_program])
    db_session.flush()
    faculty = models.User(email="faculty1@example.com", password="secret", role="faculty", campus_id=campus.id, program_id=program.id)
    second_faculty = models.User(email="faculty2@example.com", password="secret", role="faculty", campus_id=campus.id, program_id=program.id)
    other_faculty = models.User(email="faculty3@example.com", password="secret", role="faculty", campus_id=campus.id, program_id=other_program.id)
    other_campus = models.Campus(name="Other Campus", code="OTHER")
    db_session.add(other_campus)
    db_session.flush()
    mismatched_campus_faculty = models.User(email="faculty4@example.com", password="secret", role="faculty", campus_id=other_campus.id, program_id=program.id)
    db_session.add_all([faculty, second_faculty, other_faculty, mismatched_campus_faculty])
    db_session.flush()
    subject = models.Subject(name="Educational Assessment", department_id=department.id, program_id=program.id)
    unrelated_subject = models.Subject(name="Elementary Curriculum", department_id=department.id, user_id=other_faculty.id)
    department_only_subject = models.Subject(name="Department-wide Subject", department_id=department.id)
    db_session.add_all([subject, unrelated_subject, department_only_subject])
    db_session.flush()
    db_session.flush()
    first_question = models.GeneratedQuestion(subject_id=subject.id, user_id=faculty.id, question="First faculty question")
    second_question = models.GeneratedQuestion(subject_id=subject.id, user_id=second_faculty.id, question="Second faculty question")
    additional_question = models.GeneratedQuestion(subject_id=subject.id, user_id=other_faculty.id, question="Additional question in the same subject")
    mismatched_campus_question = models.GeneratedQuestion(subject_id=subject.id, user_id=mismatched_campus_faculty.id, question="Question from another campus")
    unrelated_question = models.GeneratedQuestion(subject_id=unrelated_subject.id, user_id=other_faculty.id, question="Other program question")
    department_only_question = models.GeneratedQuestion(subject_id=department_only_subject.id, user_id=faculty.id, question="Department-only question")
    db_session.add_all([first_question, second_question, additional_question, mismatched_campus_question, unrelated_question, department_only_question])
    db_session.commit()

    admin = models.User(email="admin@example.com", password="secret", role="campus_admin", campus_id=campus.id)
    super_admin = models.User(email="super_admin@example.com", password="secret", role="super_admin")
    db_session.add_all([admin, super_admin])
    db_session.commit()

    subjects = get_subjects(db=db_session, current_user=admin, program_id=program.id)
    questions = get_questions(subject_id=subject.id, program_id=program.id, db=db_session, current_user=admin)
    super_admin_subjects = get_subjects(db=db_session, current_user=super_admin, program_id=program.id)
    super_admin_questions = get_questions(subject_id=subject.id, program_id=program.id, db=db_session, current_user=super_admin)

    assert any(item["id"] == subject.id for item in subjects)
    assert any(item["name"] == "Educational Assessment" for item in subjects)
    assert next(item for item in subjects if item["id"] == subject.id)["question_count"] == 2
    assert all(item["id"] != unrelated_subject.id for item in subjects)
    assert all(item["id"] != department_only_subject.id for item in subjects)
    assert all(item["id"] != department_only_subject.id for item in super_admin_subjects)
    assert {item["id"] for item in questions} == {first_question.id, second_question.id}
    assert {item["id"] for item in super_admin_subjects} == {item["id"] for item in subjects}
    assert {item["id"] for item in super_admin_questions} == {first_question.id, second_question.id}


def test_campus_admin_can_load_user_profile_within_campus(db_session):
    campus = models.Campus(name="Profile Campus", code="PROFILE")
    db_session.add(campus)
    db_session.flush()
    department = models.Department(name="Profile Department", campus_id=campus.id)
    db_session.add(department)
    db_session.flush()
    admin = models.User(email="profile-admin@example.com", password="secret", role="campus_admin", campus_id=campus.id)
    user = models.User(email="profile-user@example.com", password="secret", role="faculty", campus_id=campus.id)
    db_session.add_all([admin, user])
    db_session.commit()
    subject_with_questions = models.Subject(name="Profile Subject", department_id=department.id)
    subject_without_questions = models.Subject(name="Empty Profile Subject", department_id=department.id)
    archived_subject = models.Subject(name="Archived Profile Subject", department_id=department.id, archived=True)
    unassigned_subject = models.Subject(name="Unassigned Profile Subject")
    db_session.add_all([subject_with_questions, subject_without_questions, archived_subject, unassigned_subject])
    db_session.flush()
    db_session.add(models.GeneratedQuestion(
        user_id=user.id,
        subject_id=subject_with_questions.id,
        question="Question text is not part of the profile response",
    ))
    upload = models.UploadedFile(
        user_id=user.id,
        subject_id=subject_with_questions.id,
        module_filename="module.pdf",
        syllabus_filename="syllabus.pdf",
        module_text="",
        syllabus_text="",
    )
    db_session.add(upload)
    db_session.flush()
    tos = models.TableOfSpecification(upload_id=upload.id, tos_data={}, total_items=1)
    db_session.add(tos)
    db_session.flush()
    db_session.add(models.GeneratedQuestion(
        tos_id=tos.id,
        subject_id=subject_with_questions.id,
        question="Legacy question with ownership recorded by its upload",
    ))
    db_session.add_all([
        models.GeneratedQuestion(user_id=user.id, subject_id=archived_subject.id, question="Archived subject question"),
        models.GeneratedQuestion(user_id=user.id, subject_id=unassigned_subject.id, question="Unassigned subject question"),
    ])
    db_session.commit()

    result = get_admin_user_overview(user.id, db_session, _admin=admin)

    assert result["user"]["id"] == user.id
    assert result["user"]["email"] == user.email
    assert [subject["id"] for subject in result["subjects"]] == [subject_with_questions.id]
    assert result["question_count"] == 4
    assert {question["question"] for question in result["subjects"][0]["questions"]} == {
        "Question text is not part of the profile response",
        "Legacy question with ownership recorded by its upload",
    }
    assert all("correct_answer" not in question for question in result["subjects"][0]["questions"])
    assert all("review_status" not in question for question in result["subjects"][0]["questions"])
    assert all("lifecycle_status" not in question for question in result["subjects"][0]["questions"])


def test_question_version_history_hides_legacy_review_metadata(db_session):
    question = models.GeneratedQuestion(question="Question with a legacy status")
    db_session.add(question)
    db_session.flush()
    db_session.add(models.QuestionVersion(
        question_id=question.id,
        snapshot={
            "question": question.question,
            "review_status": "approved",
            "lifecycle_status": "published",
        },
    ))
    db_session.commit()

    versions = get_question_versions(
        question.id,
        db_session,
        SimpleNamespace(role="super_admin"),
    )

    assert versions[0]["snapshot"] == {"question": question.question}


def test_activity_logs_include_campus_from_target_user(db_session):
    campus = models.Campus(name="Activity Campus", code="ACTIVITY")
    db_session.add(campus)
    db_session.flush()
    department = models.Department(name="Activity Department", campus_id=campus.id)
    db_session.add(department)
    db_session.flush()
    faculty = models.User(email="activity-faculty@example.com", password="secret", role="faculty", campus_id=campus.id, department=department.name)
    admin = models.User(email="activity-admin@example.com", password="secret", role="super_admin")
    db_session.add_all([faculty, admin])
    db_session.flush()
    db_session.add(models.ActivityLog(
        user_id=None,
        actor_id=admin.id,
        target_user_id=faculty.id,
        action="Department updated",
        details="Updated faculty assignment",
        type="user",
    ))
    db_session.commit()

    result = get_activity_logs(db_session, admin)

    assert result[0]["campus_id"] == campus.id
    assert result[0]["campus"] == campus.name


def test_campus_admin_updates_department_by_id_when_names_match_across_campuses(db_session):
    other_campus = models.Campus(name="Other Department Campus", code="OTHER-DEPT")
    admin_campus = models.Campus(name="Admin Department Campus", code="ADMIN-DEPT")
    db_session.add_all([other_campus, admin_campus])
    db_session.flush()
    other_department = models.Department(name="Computer Science", campus_id=other_campus.id)
    admin_department = models.Department(name="Computer Science", campus_id=admin_campus.id)
    db_session.add_all([other_department, admin_department])
    db_session.flush()
    admin = models.User(email="department-admin@example.com", password="secret", role="campus_admin", campus_id=admin_campus.id)
    user = models.User(email="department-user@example.com", password="secret", role="faculty", campus_id=admin_campus.id)
    db_session.add_all([admin, user])
    db_session.commit()

    result = update_user_department(
        UserDepartmentUpdateRequest(email=user.email, department_id=admin_department.id),
        db_session,
        admin=admin,
    )

    assert result["department"] == admin_department.name
    assert user.campus_id == admin_campus.id


def test_campus_admin_can_add_subject_name_used_by_another_program(db_session):
    campus = models.Campus(name="Shared Subjects Campus", code="SHARED-SUBJECTS")
    db_session.add(campus)
    db_session.flush()
    department = models.Department(name="Shared Subjects Department", campus_id=campus.id)
    db_session.add(department)
    db_session.flush()
    first_program = models.Program(name="First Program", department_id=department.id)
    second_program = models.Program(name="Second Program", department_id=department.id)
    db_session.add_all([first_program, second_program])
    db_session.flush()
    db_session.add_all([
        models.Subject(name="Research Methods", department_id=department.id, program_id=first_program.id),
        models.Subject(name="Research Methods", department_id=department.id, program_id=second_program.id, archived=True),
    ])
    db_session.commit()
    admin = models.User(email="shared-subjects-admin@example.com", password="secret", role="campus_admin", campus_id=campus.id)
    db_session.add(admin)
    db_session.commit()

    result = create_subject_manually(
        SubjectCreateRequest(
            name="Research Methods",
            department_id=department.id,
            program_id=second_program.id,
        ),
        db_session,
        current_user=admin,
    )

    assert result["name"] == "Research Methods"
    assert result["program_id"] == second_program.id


def test_faculty_subject_code_search_is_scoped_to_program_and_department():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    session_factory = sessionmaker(bind=engine)
    db_session = session_factory()
    previous_overrides = app.dependency_overrides.copy()
    campus = models.Campus(name="Subject Search Campus", code="SUBSEARCH")
    other_campus = models.Campus(name="Other Subject Search Campus", code="SUBSEARCH-OTHER")
    db_session.add_all([campus, other_campus])
    db_session.flush()
    department = models.Department(name="Computing Department", campus_id=campus.id)
    other_department = models.Department(name="Other Department", campus_id=campus.id)
    foreign_department = models.Department(name="Foreign Department", campus_id=other_campus.id)
    db_session.add_all([department, other_department, foreign_department])
    db_session.flush()
    program = models.Program(name="BSIT", department_id=department.id)
    other_program = models.Program(name="BSCS", department_id=department.id)
    foreign_program = models.Program(name="BSIT Foreign", department_id=foreign_department.id)
    db_session.add_all([program, other_program, foreign_program])
    db_session.flush()
    faculty = models.User(
        email="subject_search_faculty@example.com",
        password="secret",
        role="faculty",
        campus_id=campus.id,
        program_id=program.id,
        department=department.name,
    )
    db_session.add(faculty)
    db_session.add_all([
        models.Subject(name="Programming 1", code="IT101", department_id=department.id, program_id=program.id),
        models.Subject(name="Department Only", code="DEP101", department_id=department.id),
        models.Subject(name="Other Program", code="CS201", department_id=department.id, program_id=other_program.id),
        models.Subject(name="Mismatched Department", code="BAD201", department_id=other_department.id, program_id=program.id),
        models.Subject(name="Foreign Campus", code="EXT301", department_id=foreign_department.id, program_id=foreign_program.id),
        models.Subject(name="Uncoded", code=None, department_id=department.id, program_id=program.id),
        models.Subject(name="Archived", code="OLD101", department_id=department.id, program_id=program.id, archived=True),
    ])
    db_session.commit()

    results = get_faculty_program_subjects(db=db_session, current_user=faculty)

    assert [(subject["code"], subject["name"]) for subject in results] == [("IT101", "Programming 1")]
    assert results[0]["program_name"] == "BSIT"
    assert results[0]["department_name"] == department.name
    assert results[0]["has_cis"] is False

    def override_db():
        request_session = session_factory()
        try:
            yield request_session
        finally:
            request_session.close()

    app.dependency_overrides[get_db] = override_db
    app.dependency_overrides[get_current_user] = lambda: faculty
    try:
        with TestClient(app) as client:
            response = client.get("/api/faculty/subjects")
        assert response.status_code == 200
        assert [item["code"] for item in response.json()] == ["IT101"]
    finally:
        app.dependency_overrides.clear()
        app.dependency_overrides.update(previous_overrides)
        db_session.close()
        engine.dispose()


def test_faculty_upload_uses_only_the_selected_program_subject(db_session, monkeypatch):
    import main
    from routers import questions as questions_router

    campus = models.Campus(name="Upload Subject Campus", code="UPLOAD-SUBJECT")
    db_session.add(campus)
    db_session.flush()
    department = models.Department(name="Upload Computing", campus_id=campus.id)
    db_session.add(department)
    db_session.flush()
    program = models.Program(name="Upload BSIT", department_id=department.id)
    other_program = models.Program(name="Upload BSCS", department_id=department.id)
    db_session.add_all([program, other_program])
    db_session.flush()
    faculty = models.User(
        email="upload_subject_faculty@example.com",
        password="secret",
        role="faculty",
        campus_id=campus.id,
        program_id=program.id,
        department=department.name,
    )
    assigned_subject = models.Subject(
        name="Assigned Programming",
        code="UP101",
        department_id=department.id,
        program_id=program.id,
    )
    other_subject = models.Subject(
        name="Other Program Subject",
        code="CS202",
        department_id=department.id,
        program_id=other_program.id,
    )
    db_session.add_all([faculty, assigned_subject, other_subject])
    db_session.commit()

    monkeypatch.setattr(main, "extract_text", lambda _contents, _filename: "Mock Topic learning material.")
    monkeypatch.setattr(main, "detect_topics", lambda *_args, **_kwargs: {
        "course_title": "Assigned Programming",
        "course_code": "UP101",
        "topics": [{"name": "Mock Topic", "ilo": "ILO 1", "weight": 1.0}],
    })

    def upload_for_subject(subject_id):
        return asyncio.run(upload_files(
            request=SimpleNamespace(client=SimpleNamespace(host="subject-upload-test")),
            module_file=UploadFile(filename="module.docx", file=BytesIO(b"module")),
            syllabus_file=None,
            subject_id=subject_id,
            user_id=None,
            db=db_session,
            current_user=faculty,
        ))

    with pytest.raises(HTTPException) as denied:
        upload_for_subject(other_subject.id)
    assert denied.value.status_code == 404

    with pytest.raises(HTTPException) as missing_cis:
        upload_for_subject(assigned_subject.id)
    assert missing_cis.value.status_code == 409
    db_session.add(models.SubjectCIS(
        subject_id=assigned_subject.id,
        uploaded_by=faculty.id,
        filename="course-cis.docx",
        file_content=b"cis",
        extracted_text="Course Information Sheet text.",
    ))
    db_session.commit()

    monkeypatch.setattr(main, "detect_topics", lambda *_args, **_kwargs: {
        "course_title": "Unrelated Course",
        "course_code": "OTHER101",
        "topics": [{"name": "Mock Topic", "ilo": "ILO 1", "weight": 1.0}],
    })
    with pytest.raises(HTTPException) as mismatched_cis:
        upload_for_subject(assigned_subject.id)
    assert mismatched_cis.value.status_code == 422
    assert mismatched_cis.value.detail["code"] == "MATERIALS_MISMATCH"
    assert db_session.query(models.UploadedFile).count() == 0

    monkeypatch.setattr(main, "detect_topics", lambda *_args, **_kwargs: {
        "course_title": "Assigned Programming",
        "course_code": "UP101",
        "topics": [{"name": "Mock Topic", "ilo": "ILO 1", "weight": 1.0}],
    })
    result = upload_for_subject(assigned_subject.id)
    assert result["subject"]["code"] == "UP101"
    upload_record = db_session.query(models.UploadedFile).filter_by(id=result["upload_id"]).one()
    assert upload_record.subject_id == assigned_subject.id
    assert upload_record.syllabus_filename == "course-cis.docx"
    assert upload_record.syllabus_text == "Course Information Sheet text."
    questions_router.FILE_CACHE.pop(f"{result['upload_id']}_legacy_topics", None)


def test_active_question_upload_uses_selected_faculty_program_subject(db_session, monkeypatch):
    from routers import questions as questions_router

    campus = models.Campus(name="Active Upload Campus", code="ACTIVE-UPLOAD")
    db_session.add(campus)
    db_session.flush()
    department = models.Department(name="Active Upload Department", campus_id=campus.id)
    db_session.add(department)
    db_session.flush()
    program = models.Program(name="Active Upload Program", department_id=department.id)
    other_program = models.Program(name="Other Active Program", department_id=department.id)
    db_session.add_all([program, other_program])
    db_session.flush()
    faculty = models.User(
        email="active_upload_faculty@example.com",
        password="secret",
        role="faculty",
        campus_id=campus.id,
        program_id=program.id,
        department=department.name,
    )
    assigned_subject = models.Subject(
        name="Selected Course",
        code="ACTIVE101",
        department_id=department.id,
        program_id=program.id,
    )
    other_subject = models.Subject(
        name="Other Program Course",
        code="OTHER201",
        department_id=department.id,
        program_id=other_program.id,
    )
    db_session.add_all([faculty, assigned_subject, other_subject])
    db_session.commit()
    monkeypatch.setattr(questions_router, "extract_text", lambda *_args: "Test topic module material.")
    monkeypatch.setattr(questions_router, "parse_syllabus_text_with_ai", lambda *_args, **_kwargs: (
        "Selected Course",
        "ACTIVE101",
        [{"name": "Test topic", "weight": 1.0, "ilo": "ILO 1", "ilo_description": "Test outcome."}],
    ))

    def upload(subject_id):
        return asyncio.run(upload_and_analyze_syllabus(
            module_file=UploadFile(filename="module.docx", file=BytesIO(b"module")),
            syllabus_file=None,
            subject_id=subject_id,
            user_id=None,
            db=db_session,
            current_user=faculty,
        ))

    with pytest.raises(HTTPException) as missing_subject:
        upload(None)
    assert missing_subject.value.status_code == 422
    with pytest.raises(HTTPException) as foreign_subject:
        upload(other_subject.id)
    assert foreign_subject.value.status_code == 404

    with pytest.raises(HTTPException) as missing_cis:
        upload(assigned_subject.id)
    assert missing_cis.value.status_code == 409
    db_session.add(models.SubjectCIS(
        subject_id=assigned_subject.id,
        uploaded_by=faculty.id,
        filename="course-cis.docx",
        file_content=b"cis",
        extracted_text="Course Information Sheet text.",
    ))
    db_session.commit()

    monkeypatch.setattr(questions_router, "extract_text", lambda *_args: "Unrelated biology learning material.")
    with pytest.raises(HTTPException) as mismatched_material:
        upload(assigned_subject.id)
    assert mismatched_material.value.status_code == 422
    assert mismatched_material.value.detail["code"] == "MATERIALS_MISMATCH"
    assert db_session.query(models.UploadedFile).count() == 0

    monkeypatch.setattr(questions_router, "extract_text", lambda *_args: "Test topic module material.")
    result = upload(assigned_subject.id)
    assert result["subject"]["name"] == assigned_subject.name
    assert result["subject"]["code"] == assigned_subject.code
    upload_record = db_session.query(models.UploadedFile).filter_by(id=int(result["upload_id"])).one()
    assert upload_record.subject_id == assigned_subject.id
    assert upload_record.syllabus_filename == "course-cis.docx"
    assert upload_record.syllabus_text == "Course Information Sheet text."
    questions_router.FILE_CACHE.pop(f"{result['upload_id']}_metadata", None)


def test_super_admin_overview_counts_only_campus_attributed_questions(db_session):
    lipa = models.Campus(name="Alpha Campus", code="ALPHA")
    main = models.Campus(name="Beta Campus", code="BETA")
    db_session.add_all([lipa, main])
    db_session.flush()
    lipa_department = models.Department(name="Alpha Department", campus_id=lipa.id)
    main_department = models.Department(name="Beta Department", campus_id=main.id)
    db_session.add_all([lipa_department, main_department])
    db_session.flush()
    lipa_program = models.Program(name="Alpha Program", department_id=lipa_department.id)
    main_program = models.Program(name="Beta Program", department_id=main_department.id)
    db_session.add_all([lipa_program, main_program])
    db_session.flush()
    lipa_faculty = models.User(email="alpha_faculty@example.com", password="secret", role="faculty", campus_id=lipa.id, program_id=lipa_program.id)
    main_faculty = models.User(email="beta_faculty@example.com", password="secret", role="faculty", campus_id=main.id, program_id=main_program.id)
    db_session.add_all([lipa_faculty, main_faculty])
    db_session.flush()
    subject = models.Subject(name="Overview Subject", department_id=lipa_department.id)
    db_session.add(subject)
    db_session.flush()
    lipa_questions = [
        models.GeneratedQuestion(subject_id=subject.id, user_id=lipa_faculty.id, bloom_level="Remember" if index < 5 else "Analyze", question=f"Lipa question {index}")
        for index in range(10)
    ]
    main_questions = [
        models.GeneratedQuestion(subject_id=subject.id, user_id=main_faculty.id, bloom_level="Remember" if index < 5 else "Analyze", question=f"Main question {index}")
        for index in range(10)
    ]
    unassigned_questions = [
        models.GeneratedQuestion(bloom_level="Create", question=f"Unassigned question {index}")
        for index in range(3)
    ]
    db_session.add_all(lipa_questions + main_questions + unassigned_questions)
    db_session.commit()

    admin = models.User(email="overview_admin@example.com", password="secret", role="super_admin")
    db_session.add(admin)
    db_session.commit()

    overview = get_super_admin_overview(db=db_session, _admin=admin)

    campuses = {campus["name"]: campus for campus in overview["campuses"]}
    assert overview["totals"]["questions"] == 20
    assert campuses["Alpha Campus"]["questions"] == 10
    assert campuses["Beta Campus"]["questions"] == 10
    assert campuses["Alpha Campus"]["bloom_distribution"] == {"Remember": 5, "Analyze": 5}
    assert campuses["Beta Campus"]["bloom_distribution"] == {"Remember": 5, "Analyze": 5}
    assert overview["bloom_distribution"] == {"Remember": 10, "Analyze": 10}


def test_campus_admin_overview_is_limited_to_assigned_campus(db_session):
    alpha = models.Campus(name="Campus Overview Alpha", code="COV-ALPHA")
    beta = models.Campus(name="Campus Overview Beta", code="COV-BETA")
    db_session.add_all([alpha, beta])
    db_session.flush()
    alpha_department = models.Department(name="Alpha Overview Department", campus_id=alpha.id)
    alpha_other_department = models.Department(name="Alpha Other Department", campus_id=alpha.id)
    beta_department = models.Department(name="Beta Overview Department", campus_id=beta.id)
    db_session.add_all([alpha_department, alpha_other_department, beta_department])
    db_session.flush()
    alpha_program = models.Program(name="Alpha Overview Program", department_id=alpha_department.id)
    alpha_other_program = models.Program(name="Alpha Other Program", department_id=alpha_other_department.id)
    beta_program = models.Program(name="Beta Overview Program", department_id=beta_department.id)
    db_session.add_all([alpha_program, alpha_other_program, beta_program])
    db_session.flush()
    alpha_user = models.User(email="campus_overview_alpha@example.com", password="test-password", role="faculty", campus_id=alpha.id, program_id=alpha_program.id)
    alpha_other_user = models.User(email="campus_overview_other@example.com", password="test-password", role="faculty", campus_id=alpha.id, program_id=alpha_other_program.id)
    alpha_unassigned_user = models.User(email="campus_overview_unassigned@example.com", password="test-password", role="faculty", campus_id=alpha.id)
    beta_user = models.User(email="campus_overview_beta@example.com", password="test-password", role="faculty", campus_id=beta.id, program_id=beta_program.id)
    db_session.add_all([alpha_user, alpha_other_user, alpha_unassigned_user, beta_user])
    db_session.flush()
    alpha_subject = models.Subject(name="Alpha Overview Subject", department_id=alpha_department.id)
    alpha_other_subject = models.Subject(name="Alpha Other Subject", department_id=alpha_other_department.id)
    beta_subject = models.Subject(name="Beta Overview Subject", department_id=beta_department.id)
    db_session.add_all([alpha_subject, alpha_other_subject, beta_subject])
    db_session.flush()
    db_session.add_all([
        models.GeneratedQuestion(subject_id=alpha_subject.id, user_id=alpha_user.id, bloom_level="Remember", question="Alpha question", question_type="Multiple Choice"),
        models.GeneratedQuestion(subject_id=alpha_subject.id, user_id=alpha_unassigned_user.id, bloom_level="Remember", question="Alpha subject question", question_type="Multiple Choice"),
        models.GeneratedQuestion(subject_id=alpha_other_subject.id, user_id=alpha_other_user.id, bloom_level="Analyze", question="Alpha other question", question_type="Essay"),
        models.GeneratedQuestion(subject_id=beta_subject.id, user_id=beta_user.id, bloom_level="Analyze", question="Beta question", question_type="Essay"),
    ])
    db_session.commit()

    overview = get_admin_campus_overview(
        db=db_session,
        _admin=SimpleNamespace(role="campus_admin", campus_id=alpha.id),
    )

    assert overview["totals"]["campuses"] == 1
    assert overview["totals"]["departments"] == 2
    assert overview["totals"]["programs"] == 2
    assert overview["totals"]["subjects"] == 2
    assert overview["totals"]["users"] == 3
    assert overview["totals"]["faculty"] == 3
    assert overview["totals"]["questions"] == 3
    assert overview["campuses"][0]["name"] == "Campus Overview Alpha"
    assert overview["bloom_distribution"] == {"Remember": 2, "Analyze": 1}
    assert overview["question_type_distribution"] == {"Multiple Choice": 2, "Essay": 1}

    filtered = get_admin_campus_overview(
        db=db_session,
        department_id=alpha_department.id,
        _admin=SimpleNamespace(role="campus_admin", campus_id=alpha.id),
    )
    assert filtered["totals"]["departments"] == 1
    assert filtered["totals"]["programs"] == 1
    assert filtered["totals"]["subjects"] == 1
    assert filtered["totals"]["faculty"] == 1
    assert filtered["totals"]["questions"] == 2
    assert filtered["bloom_distribution"] == {"Remember": 2}
    assert filtered["question_type_distribution"] == {"Multiple Choice": 2}
    with pytest.raises(HTTPException) as error:
        get_admin_campus_overview(
            db=db_session,
            department_id=beta_department.id,
            _admin=SimpleNamespace(role="campus_admin", campus_id=alpha.id),
        )
    assert error.value.status_code == 404


def test_super_admin_ai_usage_aggregates_and_filters_by_campus_date_and_status(db_session):
    alpha = models.Campus(name="Usage Alpha", code="USAGE-ALPHA")
    beta = models.Campus(name="Usage Beta", code="USAGE-BETA")
    db_session.add_all([alpha, beta])
    db_session.flush()
    alpha_user = models.User(
        email="usage_alpha@example.com",
        password="secret",
        role="faculty",
        campus_id=alpha.id,
    )
    beta_user = models.User(
        email="usage_beta@example.com",
        password="secret",
        role="faculty",
        campus_id=beta.id,
    )
    super_admin = models.User(email="usage_admin@example.com", password="secret", role="super_admin")
    db_session.add_all([alpha_user, beta_user, super_admin])
    db_session.flush()
    db_session.add_all([
        models.AIUsage(
            user_id=alpha_user.id,
            campus_id=alpha.id,
            generated_at=datetime(2026, 9, 8, 12),
            request_type="question_generation",
            requested_question_count=8,
            generated_question_count=8,
            gemini_model="gemini-flash-lite-latest",
            gemini_api_call_count=2,
            status="success",
            request_duration_ms=2400,
            input_tokens=10,
            output_tokens=5,
            total_tokens=15,
        ),
        models.AIUsage(
            user_id=alpha_user.id,
            campus_id=alpha.id,
            generated_at=datetime(2026, 8, 20, 12),
            request_type="question_recreation",
            requested_question_count=2,
            generated_question_count=1,
            gemini_model="gemini-flash-lite-latest",
            gemini_api_call_count=3,
            status="failed",
            error_type="RuntimeError",
        ),
        models.AIUsage(
            user_id=beta_user.id,
            campus_id=beta.id,
            generated_at=datetime(2026, 9, 9, 12),
            request_type="question_generation",
            requested_question_count=4,
            generated_question_count=0,
            status="rate_limited",
            error_type="generation_cooldown",
        ),
        models.AIUsage(
            user_id=beta_user.id,
            campus_id=beta.id,
            generated_at=datetime(2026, 9, 10, 12),
            request_type="syllabus_analysis",
            requested_question_count=0,
            generated_question_count=0,
            gemini_model="gemini-flash-lite-latest",
            gemini_api_call_count=1,
            status="success",
            input_tokens=20,
            output_tokens=5,
            total_tokens=25,
        ),
    ])
    db_session.commit()

    overview = get_super_admin_ai_usage(db=db_session, _admin=super_admin)
    assert overview["totals"]["total_ai_requests"] == 4
    assert overview["totals"]["generation_requests"] == 3
    assert overview["totals"]["total_questions_generated"] == 9
    assert overview["totals"]["successful_requests"] == 2
    assert overview["totals"]["failed_requests"] == 1
    assert overview["totals"]["rate_limit_events"] == 1
    assert overview["totals"]["gemini_api_calls"] == 6
    assert overview["token_usage"]["available"] is False
    campuses = {row["campus"]: row for row in overview["requests_by_campus"]}
    assert campuses["Usage Alpha"]["requests"] == 2
    assert campuses["Usage Alpha"]["questions"] == 9
    assert campuses["Usage Beta"]["requests"] == 2
    assert campuses["Usage Beta"]["questions"] == 0

    filtered = get_super_admin_ai_usage(
        db=db_session,
        _admin=super_admin,
        campus_id=alpha.id,
        date_from=date(2026, 9, 1),
        date_to=date(2026, 9, 30),
        request_status="success",
    )
    assert filtered["totals"]["total_ai_requests"] == 1
    assert filtered["totals"]["total_questions_generated"] == 8
    assert filtered["token_usage"] == {
        "available": True,
        "input_tokens": 10,
        "output_tokens": 5,
        "total_tokens": 15,
        "requests_with_gemini_calls": 1,
    }


def test_super_admin_ai_usage_route_requires_super_admin():
    route = next(
        route for route in app.routes
        if getattr(route, "path", None) == "/api/super-admin/ai-usage"
    )
    dependency_calls = {dependency.call for dependency in route.dependant.dependencies}
    assert require_super_admin in dependency_calls
    for user in (
        SimpleNamespace(role="campus_admin", campus_id=1),
        SimpleNamespace(role="faculty", campus_id=1),
    ):
        with pytest.raises(HTTPException) as error:
            require_super_admin(user)
        assert error.value.status_code == 403


def test_campus_admin_cannot_access_foreign_campus_ai_usage_api(db_session):
    campus_admin = SimpleNamespace(role="campus_admin", campus_id=10)

    def override_db():
        yield db_session

    previous_overrides = app.dependency_overrides.copy()
    app.dependency_overrides[get_db] = override_db
    app.dependency_overrides[get_current_user] = lambda: campus_admin
    try:
        with TestClient(app) as client:
            response = client.get("/api/super-admin/ai-usage", params={"campus_id": 11})
        assert response.status_code == 403
    finally:
        app.dependency_overrides.clear()
        app.dependency_overrides.update(previous_overrides)


def test_campus_admin_ai_usage_is_forced_to_assigned_campus(db_session):
    alpha = models.Campus(name="Scoped Usage Alpha", code="SCOPED-ALPHA")
    beta = models.Campus(name="Scoped Usage Beta", code="SCOPED-BETA")
    db_session.add_all([alpha, beta])
    db_session.flush()
    alpha_user = models.User(email="scoped_usage_alpha@example.com", password="test-password", role="faculty", campus_id=alpha.id)
    beta_user = models.User(email="scoped_usage_beta@example.com", password="test-password", role="faculty", campus_id=beta.id)
    db_session.add_all([alpha_user, beta_user])
    db_session.flush()
    db_session.add_all([
        models.AIUsage(
            user_id=alpha_user.id,
            campus_id=alpha.id,
            generated_at=datetime(2026, 9, 8, 12),
            request_type="question_generation",
            requested_question_count=3,
            generated_question_count=2,
            gemini_model="alpha-model",
            gemini_api_call_count=1,
            status="success",
        ),
        models.AIUsage(
            user_id=beta_user.id,
            campus_id=beta.id,
            generated_at=datetime(2026, 9, 8, 12),
            request_type="question_generation",
            requested_question_count=10,
            generated_question_count=10,
            gemini_model="beta-model",
            gemini_api_call_count=4,
            status="success",
        ),
    ])
    db_session.commit()
    campus_admin = SimpleNamespace(role="campus_admin", campus_id=alpha.id)

    usage = get_admin_ai_usage(db=db_session, _admin=campus_admin)
    assert usage["totals"]["total_ai_requests"] == 1
    assert usage["totals"]["total_questions_generated"] == 2
    assert usage["requests_by_campus"] == [{
        "campus_id": alpha.id,
        "campus": "Scoped Usage Alpha",
        "requests": 1,
        "questions": 2,
    }]
    assert usage["available_models"] == ["alpha-model"]
    with pytest.raises(HTTPException) as error:
        get_admin_ai_usage(db=db_session, campus_id=beta.id, _admin=campus_admin)
    assert error.value.status_code == 403


def test_tos_leadership_uses_faculty_program_department_before_subject_department(db_session):
    faculty_campus = models.Campus(name="Faculty Campus", code="FACULTY-CAMPUS")
    subject_campus = models.Campus(name="Subject Campus", code="SUBJECT-CAMPUS")
    db_session.add_all([faculty_campus, subject_campus])
    db_session.flush()
    faculty_department = models.Department(name="Education Department", code="EDU", campus_id=faculty_campus.id, dean_name="Faculty Campus Dean")
    subject_department = models.Department(name="Education Department", code="OTHER-EDU", campus_id=subject_campus.id, dean_name="Subject Campus Dean")
    db_session.add_all([faculty_department, subject_department])
    db_session.flush()
    faculty_program = models.Program(name="Education Program", code="EDU-BS", department_id=faculty_department.id)
    subject_program = models.Program(name="Other Education Program", code="OTHER-BS", department_id=subject_department.id)
    db_session.add_all([faculty_program, subject_program])
    db_session.flush()
    faculty = models.User(email="faculty_tos@example.com", password="secret", role="faculty", campus_id=faculty_campus.id, program_id=faculty_program.id, department=faculty_department.name)
    subject = models.Subject(name="Cross-listed Subject", department_id=subject_department.id, program_id=subject_program.id)
    db_session.add_all([faculty, subject])
    db_session.commit()

    leadership = _resolve_department_leadership(db_session, creator=faculty, subject=subject)

    assert leadership["department_name"] == faculty_department.name
    assert leadership["department_code"] == "EDU"
    assert leadership["dean_name"] == "Faculty Campus Dean"
    assert leadership["program_code"] == "EDU-BS"


def test_question_bank_tos_export_uses_authenticated_faculty_department_dean(db_session):
    other_campus = models.Campus(name="Other TOS Campus", code="OTHER-TOS")
    faculty_campus = models.Campus(name="Faculty TOS Campus", code="FACULTY-TOS")
    db_session.add_all([other_campus, faculty_campus])
    db_session.flush()
    other_department = models.Department(name="Education Department", campus_id=other_campus.id, dean_name="Other Campus Dean")
    faculty_department = models.Department(name="Education Department", campus_id=faculty_campus.id, dean_name="Faculty Campus Dean")
    db_session.add_all([other_department, faculty_department])
    db_session.flush()
    program = models.Program(name="Teacher Education", department_id=faculty_department.id)
    db_session.add(program)
    db_session.flush()
    faculty = models.User(email="tos_faculty@example.com", password="secret", role="faculty", campus_id=faculty_campus.id, program_id=program.id, department=faculty_department.name)
    untrusted_user = models.User(email="other_user@example.com", password="secret", role="faculty", campus_id=other_campus.id)
    db_session.add_all([faculty, untrusted_user])
    db_session.flush()
    subject = models.Subject(
        name="Faculty TOS Subject",
        user_id=faculty.id,
        department_id=other_department.id,
    )
    db_session.add(subject)
    db_session.flush()
    question = models.GeneratedQuestion(
        subject_id=subject.id,
        user_id=faculty.id,
        topic_name="Foundations",
        bloom_level="Remember",
        question_type="MCQ",
        question="Which principle guides teaching?",
    )
    db_session.add(question)
    db_session.commit()

    response = export_question_bank_tos(
        subject_id=subject.id,
        question_ids=str(question.id),
        exam_type="Final Exam",
        semester="First Semester",
        academic_year="2026-2027",
        subcolumn_a_hours='{"Foundations": 1}',
        selected_topics='["Foundations"]',
        user_id=untrusted_user.id,
        db=db_session,
        current_user=faculty,
    )
    workbook = openpyxl.load_workbook(BytesIO(response.body))
    values = {str(cell.value).strip() for row in workbook.active.iter_rows() for cell in row if cell.value is not None}

    assert "Faculty Campus Dean" in values
    assert "Other Campus Dean" not in values


def test_legacy_saved_tos_export_restores_bloom_question_columns(db_session):
    faculty = models.User(
        email="legacy_tos_export@example.com",
        password="not-used",
        role="super_admin",
        name="Prof. Test",
    )
    db_session.add(faculty)
    db_session.flush()
    subject = models.Subject(name="Legacy TOS Subject", code="CS101", user_id=faculty.id)
    db_session.add(subject)
    db_session.flush()
    upload = models.UploadedFile(
        user_id=faculty.id,
        subject_id=subject.id,
        module_filename="module.pdf",
        syllabus_filename="syllabus.pdf",
        module_text="",
        syllabus_text="",
    )
    db_session.add(upload)
    db_session.flush()
    tos = models.TableOfSpecification(
        upload_id=upload.id,
        tos_data=[{
            "topic_name": "Database Design",
            "ilo": "Create logical models.",
            "hours_a": 3,
            "minutes_b": 0.6,
            "items": 2,
            "bloom_counts": {"Remember": 2},
        }],
        total_items=2,
    )
    db_session.add(tos)
    db_session.flush()
    db_session.add_all([
        models.GeneratedQuestion(
            tos_id=tos.id,
            subject_id=subject.id,
            user_id=faculty.id,
            topic_name="Database Design",
            bloom_level="Remember",
            question_type="MCQ",
            question="Question one?",
        ),
        models.GeneratedQuestion(
            tos_id=tos.id,
            subject_id=subject.id,
            user_id=faculty.id,
            topic_name="Database Design",
            bloom_level="Understand",
            question_type="MCQ",
            question="Question two?",
        ),
    ])
    db_session.commit()

    async def download_tos():
        response = await export_institutional_tos(
            upload_id=str(upload.id),
            user_id=None,
            exam_type=None,
            semester=None,
            academic_year=None,
            db=db_session,
            current_user=faculty,
        )
        return b"".join([chunk async for chunk in response.body_iterator])

    workbook = openpyxl.load_workbook(BytesIO(asyncio.run(download_tos())))
    worksheet = workbook.active

    assert worksheet["G23"].value == "1"
    assert worksheet["I23"].value == "2"
    assert worksheet["K23"].value in ("", None)


def test_department_creation_is_scoped_to_campus(db_session):
    campus_one = models.Campus(name="Campus One", code="C1")
    campus_two = models.Campus(name="Campus Two", code="C2")
    db_session.add_all([campus_one, campus_two])
    db_session.flush()

    admin = models.User(
        email="campus_admin@example.com",
        password="secret",
        role="campus_admin",
        campus_id=campus_one.id,
    )
    db_session.add(admin)
    db_session.commit()

    first = create_department(
        DepartmentCreateRequest(name="Computer Science", code="CS", campus_id=campus_one.id),
        db_session,
        admin,
    )
    assert first["campus_id"] == campus_one.id

    super_admin = models.User(
        email="super_admin@example.com",
        password="secret",
        role="super_admin",
    )
    db_session.add(super_admin)
    db_session.commit()
    second = create_department(
        DepartmentCreateRequest(name="Computer Science", code="CS", campus_id=campus_two.id),
        db_session,
        super_admin,
    )
    assert second["campus_id"] == campus_two.id

    with pytest.raises(HTTPException) as access_error:
        create_department(
            DepartmentCreateRequest(name="Data Science", code="DS", campus_id=campus_two.id),
            db_session,
            admin,
        )

    assert access_error.value.status_code == 403

    with pytest.raises(HTTPException) as error:
        create_department(
            DepartmentCreateRequest(name="Computer Science", code="CS-ALT", campus_id=campus_one.id),
            db_session,
            admin,
        )

    assert error.value.status_code == 400