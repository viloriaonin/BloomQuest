import io
import os
import sys
from types import SimpleNamespace

import pytest
from fastapi import UploadFile

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from main import normalize_email, read_upload_bytes
from routers.assessment import group_questions_by_type
from routers.questions import TOSGenerationPayload


def test_normalize_email_trims_and_lowercases():
    assert normalize_email("  User@Example.com ") == "user@example.com"


def test_normalize_email_handles_common_domain_typos():
    assert normalize_email("user@example,com") == "user@example.com"
    assert normalize_email("user@example;com") == "user@example.com"


@pytest.mark.asyncio
async def test_read_upload_bytes_rejects_image_content_even_with_allowed_extension():
    png_bytes = b"\x89PNG\r\n\x1a\n" + b"fake-image-data"
    file = UploadFile(filename="report.pdf", file=io.BytesIO(png_bytes))

    with pytest.raises(Exception):
        await read_upload_bytes(file, "module_file")


def test_tos_generation_payload_accepts_frontend_question_types():
    payload = TOSGenerationPayload(
        upload_id="demo-upload",
        total_items=10,
        whole_total_points=50,
        question_types=["MCQ", "True or False", "Identification", "Matching Type", "Enumeration", "Essay", "Situational"],
        selected_topic_indices=[0, 1],
        subcolumn_a_hours={"0": "3.0", "1": "2.0"},
    )

    assert payload.question_types == ["MCQ", "True or False", "Identification", "Matching Type", "Enumeration", "Essay", "Situational"]


def test_group_questions_by_type_keeps_same_question_types_together():
    questions = [
        SimpleNamespace(id=1, question_type="Essay", question="Explain the process."),
        SimpleNamespace(id=2, question_type="MCQ", question="Which is correct?"),
        SimpleNamespace(id=3, question_type="MCQ", question="Another MCQ?"),
        SimpleNamespace(id=4, question_type="Essay", question="Reflect on the result."),
        SimpleNamespace(id=5, question_type="True or False", question="This is true?"),
    ]

    grouped = group_questions_by_type(questions)

    assert [group["type"] for group in grouped] == ["Essay", "MCQ", "True or False"]
    assert [q.id for q in grouped[0]["questions"]] == [1, 4]
    assert [q.id for q in grouped[1]["questions"]] == [2, 3]
    assert [q.id for q in grouped[2]["questions"]] == [5]


def test_answer_key_uses_option_letters_for_mcq_and_matching_type():
    from routers.assessment import _format_answer_key_value

    mcq = SimpleNamespace(
        question_type="MCQ",
        options=["Alpha", "Bravo", "Charlie"],
        correct_answer="Charlie",
    )
    assert _format_answer_key_value(mcq) == "C"

    matching = SimpleNamespace(
        question_type="Matching Type",
        options={
            "left_items": ["A", "B"],
            "right_items": ["Alpha", "Bravo"],
        },
        correct_answer={"A": "Alpha", "B": "Bravo"},
    )
    assert _format_answer_key_value(matching) == "A -> A; B -> B"


def test_matching_type_falls_back_to_direct_left_and_right_fields():
    from routers.assessment import _get_matching_items

    question = SimpleNamespace(
        question_type="Matching Type",
        left_items=["IT concept", "Course code"],
        right_items=["System integration", "Prerequisite"],
        correct_answer={"IT concept": "System integration", "Course code": "Prerequisite"},
    )

    left_items, right_items = _get_matching_items(question)
    assert left_items == ["IT concept", "Course code"]
    assert right_items == ["System integration", "Prerequisite"]


def test_matching_type_answer_key_uses_letters_for_raw_brace_string_answers():
    from routers.assessment import _format_answer_key_value

    question = SimpleNamespace(
        question_type="Matching Type",
        options={
            "left_items": ["A", "B"],
            "right_items": ["Alpha", "Bravo"],
        },
        correct_answer='{"A": "Alpha", "B": "Bravo"}',
    )

    assert _format_answer_key_value(question) == "A -> A; B -> B"


def test_matching_type_accepts_pair_string_answers_from_db():
    from main import matching_choices

    question = SimpleNamespace(
        question_type="Matching Type",
        options={"left_items": ["A", "B"], "right_items": ["Alpha", "Bravo"]},
        correct_answer=["A -> Alpha", "B -> Bravo"],
    )

    left_items, right_items = matching_choices(question)
    assert left_items == ["A", "B"]
    assert right_items == ["Alpha", "Bravo"]


def test_assessment_docx_includes_matching_type_table_for_options_payloads():
    from docx import Document
    from routers.questions import _build_assessment_docx

    question = {
        "question": "Match the concept to its definition.",
        "question_type": "Matching Type",
        "options": {"left_items": ["Concept A", "Concept B"], "right_items": ["Definition A", "Definition B"]},
        "correct_answer": {"Concept A": "Definition A", "Concept B": "Definition B"},
    }

    content = _build_assessment_docx([question], "Course Title", "CS 101")
    document = Document(io.BytesIO(content))
    assert len(document.tables) >= 1

    table_rows = [
        [cell.text for cell in row.cells]
        for row in document.tables[0].rows
    ]

    assert ["Column A", "Column B"] in table_rows
    assert any("Concept A" in cell for row in table_rows for cell in row)
    assert any("Definition A" in cell for row in table_rows for cell in row)
    assert any("1. Concept A" in cell for row in table_rows for cell in row)
    assert any("A. Definition A" in cell for row in table_rows for cell in row)


def test_preview_saved_file_handles_merged_cells_in_spreadsheet_downloads():
    import openpyxl
    from main import preview_saved_file

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TOS"
    ws["A1"] = "TABLE OF SPECIFICATIONS"
    ws.merge_cells("A1:C1")
    ws["A2"] = "Course code"
    ws["B2"] = "CS 101"

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    class DummyLog:
        id = 999
        user_id = 1
        type = "download"
        filename = "sample-tos.xlsx"
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        file_content = stream.getvalue()

    db = SimpleNamespace(query=lambda *args, **kwargs: DummyQuery(DummyLog()))

    class DummyQuery:
        def __init__(self, log):
            self._log = log

        def filter(self, *args, **kwargs):
            return self

        def first(self):
            return self._log

    result = preview_saved_file(999, user_id=1, db=db)

    assert result["kind"] == "html"
    assert "TOS" in result["content"]
    assert "TABLE OF SPECIFICATIONS" in result["content"]
