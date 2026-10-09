import os
import sys
from io import BytesIO
from urllib.parse import parse_qs, urlparse

import pytest
from docx import Document
from fastapi import HTTPException
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import models
import main
from database import Base, get_db
from main import app, hash_password, verify_password
from security import assert_user_subject_campus_access, get_current_user


@pytest.fixture
def db_session():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    try:
        yield session
    finally:
        session.close()
        engine.dispose()


def with_test_dependencies(db_session, user):
    def override_db():
        yield db_session

    previous = app.dependency_overrides.copy()
    app.dependency_overrides[get_db] = override_db
    app.dependency_overrides[get_current_user] = lambda: user
    return previous


def test_department_admin_can_manage_only_assigned_academic_data(db_session):
    campus = models.Campus(name="Test Campus", code="TEST")
    db_session.add(campus)
    db_session.flush()
    assigned_department = models.Department(name="Assigned Department", campus_id=campus.id)
    other_department = models.Department(name="Other Department", campus_id=campus.id)
    db_session.add_all([assigned_department, other_department])
    db_session.flush()
    assigned_program = models.Program(name="Assigned Program", department_id=assigned_department.id)
    other_program = models.Program(name="Other Program", department_id=other_department.id)
    db_session.add_all([assigned_program, other_program])
    db_session.flush()
    assigned_subject = models.Subject(
        name="Assigned Subject",
        department_id=assigned_department.id,
        program_id=assigned_program.id,
    )
    other_subject = models.Subject(
        name="Other Subject",
        department_id=other_department.id,
        program_id=other_program.id,
    )
    db_session.add_all([assigned_subject, other_subject])
    department_admin = models.User(
        email="department-admin@example.com",
        password=hash_password("department-password"),
        role="department_admin",
        campus_id=campus.id,
        admin_department_id=assigned_department.id,
        archived=False,
    )
    db_session.add(department_admin)
    db_session.commit()

    assert_user_subject_campus_access(db_session, department_admin, assigned_subject)
    with pytest.raises(HTTPException) as error:
        assert_user_subject_campus_access(db_session, department_admin, other_subject)
    assert error.value.status_code == 403

    previous = with_test_dependencies(db_session, department_admin)
    try:
        with TestClient(app) as client:
            hierarchy = client.get("/api/academic-hierarchy")
            assert hierarchy.status_code == 200
            visible_departments = [
                department
                for visible_campus in hierarchy.json()["campuses"]
                for department in visible_campus["departments"]
            ]
            assert [department["id"] for department in visible_departments] == [assigned_department.id]

            subjects = client.get("/api/subjects")
            assert subjects.status_code == 200
            assert {subject["name"] for subject in subjects.json()} == {"Assigned Subject"}

            created_program = client.post("/api/programs", json={
                "name": "New Program",
                "department_id": assigned_department.id,
            })
            assert created_program.status_code == 201

            denied_program = client.post("/api/programs", json={
                "name": "Unauthorized Program",
                "department_id": other_department.id,
            })
            assert denied_program.status_code == 403

            dean_update = client.put(
                f"/api/departments/{assigned_department.id}/dean",
                json={"name": "Assigned Dean"},
            )
            assert dean_update.status_code == 200

            denied_dean_update = client.put(
                f"/api/departments/{other_department.id}/dean",
                json={"name": "Unauthorized Dean"},
            )
            assert denied_dean_update.status_code == 403

            created_subject = client.post("/api/subjects", json={
                "name": "New Subject",
                "department_id": assigned_department.id,
                "program_id": assigned_program.id,
            })
            assert created_subject.status_code == 400
            assert "Course Information Sheet" in created_subject.json()["detail"]

            denied_subject = client.post("/api/subjects", json={
                "name": "Unauthorized Subject",
                "department_id": other_department.id,
                "program_id": other_program.id,
            })
            assert denied_subject.status_code == 403

            denied_admin_view = client.get("/api/admin/insights")
            assert denied_admin_view.status_code == 403
    finally:
        app.dependency_overrides.clear()
        app.dependency_overrides.update(previous)


def test_department_admin_user_management_is_department_scoped(db_session, monkeypatch):
    campus = models.Campus(name="User Management Campus", code="UM")
    db_session.add(campus)
    db_session.flush()
    department = models.Department(name="User Management Department", campus_id=campus.id)
    other_department = models.Department(name="Other User Management Department", campus_id=campus.id)
    db_session.add_all([department, other_department])
    db_session.flush()
    program = models.Program(name="User Management Program", department_id=department.id)
    other_program = models.Program(name="Other User Management Program", department_id=other_department.id)
    db_session.add_all([program, other_program])
    db_session.flush()
    admin = models.User(
        email="user-management-admin@example.com",
        password=hash_password("admin-password"),
        role="department_admin",
        campus_id=campus.id,
        admin_department_id=department.id,
        archived=False,
    )
    faculty = models.User(
        email="managed-faculty@example.com",
        password=hash_password("faculty-password"),
        role="faculty",
        name="Managed Faculty",
        department=department.name,
        campus_id=campus.id,
        program_id=program.id,
        archived=False,
    )
    other_faculty = models.User(
        email="other-faculty@example.com",
        password=hash_password("faculty-password"),
        role="faculty",
        name="Other Faculty",
        department=other_department.name,
        campus_id=campus.id,
        program_id=other_program.id,
        archived=False,
    )
    student = models.User(
        email="student@example.com",
        password=hash_password("student-password"),
        role="student",
        name="Student",
        department=department.name,
        campus_id=campus.id,
        archived=False,
    )
    db_session.add_all([admin, faculty, other_faculty, student])
    db_session.flush()
    subject = models.Subject(
        name="Managed Subject",
        department_id=department.id,
        program_id=program.id,
        user_id=faculty.id,
        archived=False,
    )
    db_session.add(subject)
    db_session.add(models.ActivityLog(
        user_id=faculty.id,
        action="Faculty signed in",
        type="login",
        status="success",
    ))
    db_session.commit()

    previous = with_test_dependencies(db_session, admin)
    sent_credentials = []
    monkeypatch.setattr(main, "require_email_delivery_configured", lambda: None)
    monkeypatch.setattr(
        main,
        "send_department_faculty_access_email",
        lambda *args: sent_credentials.append(args) or True,
    )
    try:
        with TestClient(app) as client:
            listed = client.get("/api/department-admin/users")
            assert listed.status_code == 200
            assert [user["email"] for user in listed.json()["users"]] == [faculty.email]

            detail = client.get(f"/api/department-admin/users/{faculty.id}")
            assert detail.status_code == 200
            assert detail.json()["subjects"][0]["name"] == "Managed Subject"
            assert detail.json()["activity"][0]["action"] == "Faculty signed in"

            changed_program = client.patch(
                f"/api/department-admin/users/{faculty.id}/program",
                json={"program_id": program.id},
            )
            assert changed_program.status_code == 200

            reset_email = client.post(
                f"/api/department-admin/users/{faculty.id}/credential-email",
                json={"action": "reset_password"},
            )
            assert reset_email.status_code == 200
            assert reset_email.json()["email_status"] == "sent"
            reset_message = sent_credentials[-1]
            assert reset_message[0] == faculty.email
            assert "set-password?token=" in reset_message[4]
            assert db_session.refresh(faculty) is None
            first_token_hash = faculty.password_setup_token_hash
            assert first_token_hash
            assert verify_password(reset_message[1], faculty.password)

            previous_password = faculty.password
            previous_token_hash = faculty.password_setup_token_hash
            monkeypatch.setattr(main, "send_department_faculty_access_email", lambda *args: False)
            failed_email = client.post(
                f"/api/department-admin/users/{faculty.id}/credential-email",
                json={"action": "reset_password"},
            )
            assert failed_email.status_code == 503
            db_session.refresh(faculty)
            assert faculty.password == previous_password
            assert faculty.password_setup_token_hash == previous_token_hash

            denied_program = client.patch(
                f"/api/department-admin/users/{faculty.id}/program",
                json={"program_id": other_program.id},
            )
            assert denied_program.status_code == 404

            archived = client.patch(
                f"/api/department-admin/users/{faculty.id}/status",
                json={"archived": True},
            )
            assert archived.status_code == 200
            assert client.get(f"/api/department-admin/users/{faculty.id}").json()["archived"] is True

            assert client.get(f"/api/department-admin/users/{other_faculty.id}").status_code == 404
            assert client.patch(
                f"/api/department-admin/users/{other_faculty.id}/status",
                json={"archived": True},
            ).status_code == 404
            assert client.get(f"/api/department-admin/users/{student.id}").status_code == 404

            deleted = client.delete(f"/api/department-admin/users/{faculty.id}")
            assert deleted.status_code == 200
            assert db_session.query(models.User).filter_by(id=faculty.id).first() is None
    finally:
        app.dependency_overrides.clear()
        app.dependency_overrides.update(previous)


def test_department_admin_can_review_only_change_requests_from_their_department(db_session):
    campus = models.Campus(name="Requests Campus", code="REQ")
    db_session.add(campus)
    db_session.flush()
    department = models.Department(name="Requests Department", campus_id=campus.id)
    other_department = models.Department(name="Other Requests Department", campus_id=campus.id)
    db_session.add_all([department, other_department])
    db_session.flush()
    department_program = models.Program(name="Department Program", department_id=department.id)
    next_department_program = models.Program(name="Next Department Program", department_id=department.id)
    other_program = models.Program(name="Other Program", department_id=other_department.id)
    db_session.add_all([department_program, next_department_program, other_program])
    db_session.flush()
    admin = models.User(
        email="requests-admin@example.com",
        password="unused",
        role="department_admin",
        campus_id=campus.id,
        admin_department_id=department.id,
        archived=False,
    )
    local_faculty = models.User(
        email="local-faculty@example.com",
        password="unused",
        role="faculty",
        campus_id=campus.id,
        department=department.name,
        program_id=department_program.id,
        archived=False,
    )
    other_faculty = models.User(
        email="other-faculty@example.com",
        password="unused",
        role="faculty",
        campus_id=campus.id,
        department=other_department.name,
        program_id=other_program.id,
        archived=False,
    )
    db_session.add_all([admin, local_faculty, other_faculty])
    db_session.flush()
    local_program_request = models.UserChangeRequest(
        user_id=local_faculty.id,
        request_type="program",
        current_value=department_program.name,
        requested_value=str(next_department_program.id),
        status="pending",
    )
    transfer_request = models.UserChangeRequest(
        user_id=local_faculty.id,
        request_type="program",
        current_value=department_program.name,
        requested_value=str(other_program.id),
        status="pending",
    )
    other_department_request = models.UserChangeRequest(
        user_id=other_faculty.id,
        request_type="program",
        current_value=other_program.name,
        requested_value=str(other_program.id),
        status="pending",
    )
    db_session.add_all([local_program_request, transfer_request, other_department_request])
    db_session.commit()

    previous = with_test_dependencies(db_session, admin)
    try:
        with TestClient(app) as client:
            listed = client.get("/api/admin/user-change-requests?status=all")
            assert listed.status_code == 200
            assert {request["id"] for request in listed.json()} == {
                local_program_request.id,
                transfer_request.id,
            }

            approved = client.patch(
                f"/api/admin/user-change-requests/{local_program_request.id}",
                json={"action": "approve"},
            )
            assert approved.status_code == 200
            assert local_faculty.program_id == next_department_program.id

            denied_transfer = client.patch(
                f"/api/admin/user-change-requests/{transfer_request.id}",
                json={"action": "approve"},
            )
            assert denied_transfer.status_code == 403

            denied_other_department = client.patch(
                f"/api/admin/user-change-requests/{other_department_request.id}",
                json={"action": "approve"},
            )
            assert denied_other_department.status_code == 404
    finally:
        app.dependency_overrides.clear()
        app.dependency_overrides.update(previous)


def test_department_admin_creates_subject_with_required_cis(db_session, monkeypatch):
    campus = models.Campus(name="CIS Campus", code="CIS")
    db_session.add(campus)
    db_session.flush()
    department = models.Department(name="CIS Department", campus_id=campus.id)
    db_session.add(department)
    db_session.flush()
    program = models.Program(name="CIS Program", department_id=department.id)
    admin = models.User(
        email="cis-admin@example.com",
        password=hash_password("cis-admin-password"),
        role="department_admin",
        campus_id=campus.id,
        admin_department_id=department.id,
        archived=False,
    )
    db_session.add_all([program, admin])
    db_session.commit()
    monkeypatch.setattr("main.log_activity", lambda *args, **kwargs: None)

    document = Document()
    document.add_paragraph("Course Information Sheet: Introduction to Computing")
    document_bytes = BytesIO()
    document.save(document_bytes)
    previous = with_test_dependencies(db_session, admin)
    try:
        with TestClient(app) as client:
            missing_cis = client.post(
                "/api/subjects/with-cis",
                data={
                    "name": "Introduction to Computing",
                    "code": "CIS101",
                    "department_id": str(department.id),
                    "program_id": str(program.id),
                },
            )
            assert missing_cis.status_code == 422
            rejected_file = client.post(
                "/api/subjects/with-cis",
                data={
                    "name": "Introduction to Computing",
                    "code": "CIS101",
                    "department_id": str(department.id),
                    "program_id": str(program.id),
                },
                files={"cis_file": ("CIS101.txt", document_bytes.getvalue(), "text/plain")},
            )
            assert rejected_file.status_code == 400
            assert db_session.query(models.Subject).filter_by(code="CIS101").first() is None
            response = client.post(
                "/api/subjects/with-cis",
                data={
                    "name": "Introduction to Computing",
                    "code": "CIS101",
                    "department_id": str(department.id),
                    "program_id": str(program.id),
                },
                files={"cis_file": ("CIS101.docx", document_bytes.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
        assert response.status_code == 201, response.text
        result = response.json()
        assert result["has_cis"] is True
        assert result["cis_filename"] == "CIS101.docx"
        subject = db_session.query(models.Subject).filter_by(id=result["id"]).one()
        cis = db_session.query(models.SubjectCIS).filter_by(subject_id=subject.id).one()
        assert cis.uploaded_by == admin.id
        assert "Introduction to Computing" in cis.extracted_text
        assert cis.file_content == document_bytes.getvalue()
        other_department = models.Department(
            name="Other Department",
            campus_id=department.campus_id,
        )
        db_session.add(other_department)
        db_session.flush()
        other_subject = models.Subject(
            name="Other Subject",
            code="OTHER101",
            department_id=other_department.id,
        )
        db_session.add(other_subject)
        db_session.commit()
        with TestClient(app) as client:
            cis_details = client.get(f"/api/subjects/{subject.id}/cis")
            cis_file = client.get(f"/api/subjects/{subject.id}/cis/file")
            cross_department_cis = client.get(f"/api/subjects/{other_subject.id}/cis")
        assert cis_details.status_code == 200, cis_details.text
        assert cis_details.json()["filename"] == "CIS101.docx"
        assert "Introduction to Computing" in cis_details.json()["extracted_text"]
        assert cis_file.status_code == 200
        assert cis_file.content == document_bytes.getvalue()
        assert cis_file.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        assert cross_department_cis.status_code == 403, cross_department_cis.text

        replacement_document = Document()
        replacement_document.add_paragraph("Updated Course Information Sheet")
        replacement_bytes = BytesIO()
        replacement_document.save(replacement_bytes)
        with TestClient(app) as client:
            replacement = client.put(
                f"/api/subjects/{subject.id}/cis",
                files={"cis_file": ("CIS101-updated.docx", replacement_bytes.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
        assert replacement.status_code == 200, replacement.text
        db_session.refresh(cis)
        assert cis.filename == "CIS101-updated.docx"
        assert "Updated Course Information Sheet" in cis.extracted_text
    finally:
        app.dependency_overrides.clear()
        app.dependency_overrides.update(previous)


def test_campus_admin_creates_faculty_for_programs_in_assigned_campus(db_session, monkeypatch):
    campus = models.Campus(name="Campus Faculty Campus", code="CFC")
    other_campus = models.Campus(name="Other Faculty Campus", code="OFC")
    db_session.add_all([campus, other_campus])
    db_session.flush()
    department = models.Department(name="Computing", campus_id=campus.id)
    other_department = models.Department(name="Other Computing", campus_id=other_campus.id)
    db_session.add_all([department, other_department])
    db_session.flush()
    program = models.Program(name="Information Technology", department_id=department.id)
    other_program = models.Program(name="Other Program", department_id=other_department.id)
    admin = models.User(
        email="campus-faculty-admin@example.edu",
        password=hash_password("CampusAdminPassword1!"),
        role="campus_admin",
        campus_id=campus.id,
        archived=False,
    )
    db_session.add_all([program, other_program, admin])
    db_session.commit()

    email_calls = []
    monkeypatch.setattr("main.require_email_delivery_configured", lambda: None)
    monkeypatch.setattr("main.send_approval_email", lambda *args: email_calls.append(args) or True)
    monkeypatch.setattr("main.log_activity", lambda *args, **kwargs: None)
    previous = with_test_dependencies(db_session, admin)
    try:
        with TestClient(app) as client:
            response = client.post(
                "/api/campus-admin/faculty-accounts",
                json={
                    "full_name": "Campus Faculty",
                    "email": "campus.faculty@example.edu",
                    "faculty_number": "000542",
                    "program_id": program.id,
                },
            )
            invalid_number = client.post(
                "/api/campus-admin/faculty-accounts",
                json={
                    "full_name": "Invalid Faculty Number",
                    "email": "invalid.faculty@example.edu",
                    "faculty_number": "54A2",
                    "program_id": program.id,
                },
            )
            cross_campus = client.post(
                "/api/campus-admin/faculty-accounts",
                json={
                    "full_name": "Out of Campus Faculty",
                    "email": "outside.faculty@example.edu",
                    "faculty_number": "000543",
                    "program_id": other_program.id,
                },
            )
        assert response.status_code == 201, response.text
        assert response.json()["faculty_number"] == "000542"
        assert response.json()["email_status"] == "sent"
        assert invalid_number.status_code == 422
        assert cross_campus.status_code == 403
        created = db_session.query(models.User).filter_by(
            email="campus.faculty@example.edu"
        ).one()
        assert created.role == "faculty"
        assert created.campus_id == campus.id
        assert created.department == department.name
        assert created.program_id == program.id
        assert created.faculty_number == "000542"
        assert len(email_calls) == 1
        assert email_calls[0][0] == created.email
        assert email_calls[0][3] == department.name
    finally:
        app.dependency_overrides.clear()
        app.dependency_overrides.update(previous)


def test_super_admin_can_create_department_admin_with_scoped_department(db_session, monkeypatch):
    campus = models.Campus(name="Admin Campus", code="ADMIN")
    db_session.add(campus)
    db_session.flush()
    department = models.Department(name="Managed Department", campus_id=campus.id)
    super_admin = models.User(
        email="super-admin@example.com",
        password=hash_password("super-password"),
        role="super_admin",
        archived=False,
    )
    db_session.add_all([department, super_admin])
    db_session.commit()
    monkeypatch.setattr("main.log_activity", lambda *args, **kwargs: None)

    previous = with_test_dependencies(db_session, super_admin)
    try:
        with TestClient(app) as client:
            response = client.post("/api/super-admin/department-admins", json={
                "name": "Department Admin",
                "email": "new-admin@example.com",
                "password": "new-admin-password",
                "department_id": department.id,
            })
            assert response.status_code == 201
            created_user = db_session.query(models.User).filter_by(
                email="new-admin@example.com"
            ).one()
            assert created_user.role == "department_admin"
            assert created_user.admin_department_id == department.id
            assert created_user.campus_id == campus.id
            assert verify_password("new-admin-password", created_user.password)
            assert response.json()["department"] == department.name
    finally:
        app.dependency_overrides.clear()
        app.dependency_overrides.update(previous)


def test_department_admin_manages_department_faculty_and_submits_requests(db_session, monkeypatch):
    campus = models.Campus(name="Faculty Campus", code="FAC")
    db_session.add(campus)
    db_session.flush()
    department = models.Department(name="Managed Department", campus_id=campus.id)
    other_department = models.Department(name="Other Department", campus_id=campus.id)
    db_session.add_all([department, other_department])
    db_session.flush()
    program = models.Program(name="Managed Program", department_id=department.id)
    other_program = models.Program(name="Other Program", department_id=other_department.id)
    db_session.add_all([program, other_program])
    db_session.flush()
    faculty = models.User(
        email="faculty@example.com",
        name="Faculty Member",
        role="faculty",
        department=department.name,
        campus_id=campus.id,
        archived=False,
    )
    other_faculty = models.User(
        email="other-faculty@example.com",
        name="Other Faculty",
        role="faculty",
        department=other_department.name,
        campus_id=campus.id,
        archived=False,
    )
    department_admin = models.User(
        email="department-admin@example.com",
        role="department_admin",
        campus_id=campus.id,
        admin_department_id=department.id,
        archived=False,
    )
    db_session.add_all([faculty, other_faculty, department_admin])
    db_session.flush()
    subject = models.Subject(
        name="Managed Subject",
        department_id=department.id,
        program_id=program.id,
    )
    db_session.add(subject)
    db_session.commit()
    monkeypatch.setattr("main.log_activity", lambda *args, **kwargs: None)

    previous = with_test_dependencies(db_session, department_admin)
    try:
        with TestClient(app) as client:
            profile = client.put(
                "/api/user/profile",
                json={"full_name": "Department Administrator"},
            )
            assert profile.status_code == 200
            assert profile.json()["full_name"] == "Department Administrator"
            assert profile.json()["department"] == department.name

            denied_profile_assignment = client.put(
                "/api/user/profile",
                json={"full_name": "Department Administrator", "program_id": program.id},
            )
            assert denied_profile_assignment.status_code == 403

            assignment = client.put(
                f"/api/faculty/{faculty.id}/program",
                json={"program_id": program.id},
            )
            assert assignment.status_code == 200
            assert assignment.json()["program_id"] == program.id

            chair_assignment = client.put(
                f"/api/programs/{program.id}/chair",
                json={"faculty_id": faculty.id},
            )
            assert chair_assignment.status_code == 200
            assert chair_assignment.json()["chair_id"] == faculty.id

            denied_assignment = client.put(
                f"/api/faculty/{other_faculty.id}/program",
                json={"program_id": program.id},
            )
            assert denied_assignment.status_code == 403

            subject_assignment = client.put(
                f"/api/subjects/{subject.id}/faculty",
                json={"faculty_id": faculty.id},
            )
            assert subject_assignment.status_code == 200
            assert subject_assignment.json()["faculty_id"] == faculty.id

            denied_subject_assignment = client.put(
                f"/api/subjects/{subject.id}/faculty",
                json={"faculty_id": other_faculty.id},
            )
            assert denied_subject_assignment.status_code == 400

            moved_subject = client.put(
                f"/api/subjects/{subject.id}",
                json={
                    "name": "Managed Subject",
                    "department_id": other_department.id,
                    "program_id": other_program.id,
                },
            )
            assert moved_subject.status_code == 403
            db_session.refresh(subject)
            assert subject.department_id == department.id
            assert subject.program_id == program.id

            details = client.put(
                f"/api/departments/{department.id}/details",
                json={"name": "Updated Department", "code": "UPDATED"},
            )
            assert details.status_code == 200, details.text
            assert faculty.department == "Updated Department"

            invitation = client.post(
                "/api/department-admin/faculty-requests",
                json={
                    "full_name": "New Faculty",
                    "email": "new-faculty@example.com",
                    "program_id": program.id,
                },
            )
            assert invitation.status_code == 201
            account_request = db_session.query(models.AccountRequest).filter_by(
                email="new-faculty@example.com"
            ).one()
            assert account_request.status == "pending"
            assert account_request.department == "Updated Department"

            change_request = client.post(
                "/api/department-academic-change-requests",
                json={
                    "title": "Coordinate shared subject",
                    "details": "Please review the shared curriculum alignment.",
                    "related_department": other_department.name,
                },
            )
            assert change_request.status_code == 201
            assert client.get("/api/department-academic-change-requests").json()[0]["title"] == "Coordinate shared subject"

            denied_review = client.put(
                f"/api/department-academic-change-requests/{change_request.json()['id']}",
                json={"status": "approved"},
            )
            assert denied_review.status_code == 403
    finally:
        app.dependency_overrides.clear()
        app.dependency_overrides.update(previous)


def test_department_admin_manages_scoped_faculty_accounts(db_session, monkeypatch):
    campus = models.Campus(name="Accounts Campus", code="ACCOUNTS")
    db_session.add(campus)
    db_session.flush()
    department = models.Department(name="Accounts Department", campus_id=campus.id)
    other_department = models.Department(name="Other Accounts Department", campus_id=campus.id)
    db_session.add_all([department, other_department])
    db_session.flush()
    program = models.Program(name="Accounts Program", department_id=department.id)
    other_program = models.Program(name="Other Accounts Program", department_id=other_department.id)
    db_session.add_all([program, other_program])
    db_session.flush()
    admin = models.User(
        email="accounts-admin@example.com",
        role="department_admin",
        campus_id=campus.id,
        admin_department_id=department.id,
        archived=False,
    )
    pending = models.AccountRequest(
        full_name="Pending Faculty",
        email="pending-faculty@example.edu",
        department=department.name,
        campus_id=campus.id,
        program_id=program.id,
        status="pending",
    )
    other_pending = models.AccountRequest(
        full_name="Other Pending Faculty",
        email="other-pending@example.edu",
        department=other_department.name,
        campus_id=campus.id,
        program_id=other_program.id,
        status="pending",
    )
    active_faculty = models.User(
        email="active-faculty@example.edu",
        password=hash_password("faculty-password"),
        role="faculty",
        name="Active Faculty",
        department=department.name,
        campus_id=campus.id,
        program_id=program.id,
        archived=False,
    )
    archived_faculty = models.User(
        email="archived-faculty@example.edu",
        password=hash_password("faculty-password"),
        role="faculty",
        name="Archived Faculty",
        department=department.name,
        campus_id=campus.id,
        program_id=program.id,
        archived=True,
    )
    db_session.add_all([
        admin,
        pending,
        other_pending,
        active_faculty,
        archived_faculty,
    ])
    db_session.commit()
    emailed = []
    monkeypatch.setattr("main.generate_temporary_password", lambda: "FacultyPassword1!")
    monkeypatch.setattr("main.send_approval_email", lambda *args: emailed.append(args))
    monkeypatch.setattr("main.log_activity", lambda *args, **kwargs: None)

    previous = with_test_dependencies(db_session, admin)
    try:
        with TestClient(app) as client:
            response = client.get("/api/department-admin/faculty-accounts")
            assert response.status_code == 200, response.text
            accounts = response.json()
            assert [item["email"] for item in accounts["pending"]] == ["pending-faculty@example.edu"]
            assert [item["email"] for item in accounts["active"]] == ["active-faculty@example.edu"]
            assert [item["email"] for item in accounts["archived"]] == ["archived-faculty@example.edu"]
            assert accounts["pending"][0]["program"] == program.name
            assert accounts["pending"][0]["requested_at"]

            approved = client.post(
                "/api/department-admin/faculty-accounts/approve",
                json={"email": pending.email},
            )
            assert approved.status_code == 200, approved.text
            created = db_session.query(models.User).filter_by(
                email=pending.email
            ).one()
            assert created.role == "faculty"
            assert created.program_id == program.id
            assert created.department == department.name
            assert verify_password("FacultyPassword1!", created.password)
            assert emailed == [(
                pending.email,
                "FacultyPassword1!",
                "Pending Faculty",
                department.name,
            )]

            archived = client.post(
                "/api/department-admin/faculty-accounts/archive",
                json={"email": active_faculty.email},
            )
            assert archived.status_code == 200
            db_session.refresh(active_faculty)
            assert active_faculty.archived is True

            restored = client.post(
                "/api/department-admin/faculty-accounts/restore",
                json={"email": archived_faculty.email},
            )
            assert restored.status_code == 200
            db_session.refresh(archived_faculty)
            assert archived_faculty.archived is False

            denied = client.post(
                "/api/department-admin/faculty-accounts/decline",
                json={"email": other_pending.email},
            )
            assert denied.status_code == 404

            declined = client.post(
                "/api/department-admin/faculty-accounts/decline",
                json={"email": "pending-faculty@example.edu"},
            )
            assert declined.status_code == 404
    finally:
        app.dependency_overrides.clear()
        app.dependency_overrides.update(previous)


def test_department_admin_manually_creates_faculty_account_and_emails_credentials(db_session, monkeypatch):
    campus = models.Campus(name="Faculty Campus", code="FACULTY")
    db_session.add(campus)
    db_session.flush()
    department = models.Department(name="Faculty Department", campus_id=campus.id)
    other_department = models.Department(name="Other Faculty Department", campus_id=campus.id)
    db_session.add_all([department, other_department])
    db_session.flush()
    program = models.Program(name="Faculty Program", department_id=department.id)
    other_program = models.Program(name="Other Program", department_id=other_department.id)
    admin = models.User(
        email="faculty-admin@example.com",
        password=hash_password("admin-password"),
        role="department_admin",
        campus_id=campus.id,
        admin_department_id=department.id,
        archived=False,
    )
    db_session.add_all([program, other_program, admin])
    db_session.commit()
    emailed = []
    email_succeeds = [True]
    monkeypatch.setattr("main.require_email_delivery_configured", lambda: None)
    monkeypatch.setattr("main.generate_temporary_password", lambda: "FacultyPassword1!")

    def capture_email(*args):
        emailed.append(args)
        return email_succeeds[0]

    monkeypatch.setattr("main.send_approval_email", capture_email)
    monkeypatch.setattr("main.log_activity", lambda *args, **kwargs: None)

    previous = with_test_dependencies(db_session, admin)
    try:
        with TestClient(app) as client:
            invalid_number = client.post(
                "/api/department-admin/faculty-accounts",
                json={
                    "full_name": "Invalid Number Faculty",
                    "email": "invalid-number@example.edu",
                    "faculty_number": "12A",
                    "program_id": program.id,
                },
            )
            assert invalid_number.status_code == 422

            response = client.post(
                "/api/department-admin/faculty-accounts",
                json={
                    "full_name": "New Faculty",
                    "email": "new-faculty@example.edu",
                    "faculty_number": "000345",
                    "program_id": program.id,
                },
            )
            assert response.status_code == 201, response.text
            assert response.json()["email"] == "new-faculty@example.edu"
            assert response.json()["faculty_number"] == "000345"
            assert response.json()["employee_id"] == "000345"
            assert response.json()["email_status"] == "sent"
            created = db_session.query(models.User).filter_by(
                email="new-faculty@example.edu"
            ).one()
            assert created.role == "faculty"
            assert created.archived is False
            assert created.program_id == program.id
            assert created.department == department.name
            assert created.faculty_number == "000345"
            listed_accounts = client.get("/api/department-admin/users")
            assert listed_accounts.status_code == 200
            assert listed_accounts.json()["users"][0]["faculty_number"] == "000345"
            assert listed_accounts.json()["users"][0]["employee_id"] == "000345"
            assert verify_password("FacultyPassword1!", created.password)
            assert len(emailed) == 1
            assert emailed[0][:4] == (
                "new-faculty@example.edu",
                "FacultyPassword1!",
                "New Faculty",
                department.name,
            )
            setup_url = emailed[0][4]
            assert setup_url.startswith("http://localhost:3000/set-password?token=")
            setup_token = parse_qs(urlparse(setup_url).query)["token"][0]
            assert created.password_setup_token_hash
            assert created.password_setup_token_expires_at

            setup_payload = {
                "token": setup_token,
                "temporary_password": "FacultyPassword1!",
                "new_password": "PermanentPassword2!",
            }
            invalid_temporary_password = client.post(
                "/api/set-initial-password",
                json={**setup_payload, "temporary_password": "WrongPassword1!"},
            )
            assert invalid_temporary_password.status_code == 400

            password_setup = client.post("/api/set-initial-password", json=setup_payload)
            assert password_setup.status_code == 200, password_setup.text
            db_session.refresh(created)
            assert verify_password("PermanentPassword2!", created.password)
            assert not created.password_setup_token_hash
            assert not created.password_setup_token_expires_at

            reused_link = client.post("/api/set-initial-password", json=setup_payload)
            assert reused_link.status_code == 400

            email_succeeds[0] = False
            failed_email = client.post(
                "/api/department-admin/faculty-accounts",
                json={
                    "full_name": "Email Failure Faculty",
                    "email": "email-failure@example.edu",
                    "program_id": program.id,
                },
            )
            assert failed_email.status_code == 201
            assert failed_email.json()["email_status"] == "failed"
            assert db_session.query(models.User).filter_by(
                email="email-failure@example.edu"
            ).one()

            duplicate = client.post(
                "/api/department-admin/faculty-accounts",
                json={
                    "full_name": "Duplicate Faculty",
                    "email": "new-faculty@example.edu",
                    "program_id": program.id,
                },
            )
            assert duplicate.status_code == 409

            out_of_scope = client.post(
                "/api/department-admin/faculty-accounts",
                json={
                    "full_name": "Other Faculty",
                    "email": "other-faculty@example.edu",
                    "program_id": other_program.id,
                },
            )
            assert out_of_scope.status_code == 404
    finally:
        app.dependency_overrides.clear()
        app.dependency_overrides.update(previous)


def test_campus_admin_creates_department_admin_dean_within_assigned_campus(db_session, monkeypatch):
    campus = models.Campus(name="Dean Campus", code="DEAN-CAMPUS")
    other_campus = models.Campus(name="Other Dean Campus", code="OTHER-DEAN")
    db_session.add_all([campus, other_campus])
    db_session.flush()
    department = models.Department(name="Dean Department", campus_id=campus.id)
    other_department = models.Department(name="Other Dean Department", campus_id=other_campus.id)
    db_session.add_all([department, other_department])
    db_session.flush()
    admin = models.User(
        email="dean-campus-admin@example.edu",
        password=hash_password("CampusAdminPassword1!"),
        role="campus_admin",
        campus_id=campus.id,
        archived=False,
    )
    db_session.add(admin)
    db_session.commit()

    email_calls = []
    monkeypatch.setattr(main, "require_email_delivery_configured", lambda: None)
    monkeypatch.setattr(main, "generate_temporary_password", lambda: "CampusDeanPassword1!")
    monkeypatch.setattr(main, "send_approval_email", lambda *args: email_calls.append(args) or True)
    monkeypatch.setattr(main, "log_activity", lambda *args, **kwargs: None)
    previous = with_test_dependencies(db_session, admin)
    try:
        with TestClient(app) as client:
            created = client.post(
                "/api/campus-admin/department-admins",
                json={
                    "full_name": "Campus Dean",
                    "email": "campus.dean@example.edu",
                    "department_id": department.id,
                },
            )
            cross_campus = client.post(
                "/api/campus-admin/department-admins",
                json={
                    "full_name": "Outside Dean",
                    "email": "outside.dean@example.edu",
                    "department_id": other_department.id,
                },
            )

        assert created.status_code == 201, created.text
        assert created.json()["role"] == "department_admin"
        assert created.json()["department"] == department.name
        assert created.json()["email_status"] == "sent"
        assert cross_campus.status_code == 403
        dean = db_session.query(models.User).filter_by(
            email="campus.dean@example.edu"
        ).one()
        db_session.refresh(department)
        assert dean.role == "department_admin"
        assert dean.admin_department_id == department.id
        assert dean.campus_id == campus.id
        assert department.dean_id == dean.id
        assert department.dean_name == dean.name
        assert verify_password("CampusDeanPassword1!", dean.password)
        assert dean.password_setup_token_hash
        assert len(email_calls) == 1
        assert email_calls[0][0] == dean.email
        assert "set-password?token=" in email_calls[0][4]
    finally:
        app.dependency_overrides.clear()
        app.dependency_overrides.update(previous)


def test_campus_admin_can_create_dean_in_demo_mode_without_email_delivery(db_session, monkeypatch):
    campus = models.Campus(name="Demo Dean Campus", code="DEMO-DEAN")
    db_session.add(campus)
    db_session.flush()
    department = models.Department(name="Demo Dean Department", campus_id=campus.id)
    admin = models.User(
        email="demo-dean-admin@example.edu",
        password=hash_password("CampusAdminPassword1!"),
        role="campus_admin",
        campus_id=campus.id,
        archived=False,
    )
    db_session.add_all([department, admin])
    db_session.commit()

    monkeypatch.setattr(main, "DEMO_EMAIL_VERIFICATION", True)
    monkeypatch.setattr(
        main,
        "require_email_delivery_configured",
        lambda: pytest.fail("Email configuration must not be required in demo mode."),
    )
    monkeypatch.setattr(main, "generate_temporary_password", lambda: "DemoDeanPassword1!")
    monkeypatch.setattr(
        main,
        "send_approval_email",
        lambda *args: pytest.fail("Demo mode must not attempt to send email."),
    )
    monkeypatch.setattr(main, "log_activity", lambda *args, **kwargs: None)
    previous = with_test_dependencies(db_session, admin)
    try:
        with TestClient(app) as client:
            response = client.post(
                "/api/campus-admin/department-admins",
                json={
                    "full_name": "Demo Dean",
                    "email": "demo.dean@example.edu",
                    "department_id": department.id,
                },
            )

            assert response.status_code == 201, response.text
            result = response.json()
            assert result["email_status"] == "demo"
            assert result["demo_temporary_password"] == "DemoDeanPassword1!"
            setup_url = result["demo_setup_url"]
            setup_token = parse_qs(urlparse(setup_url).query)["token"][0]
            assert setup_url.endswith(f"/set-password?token={setup_token}")

            setup = client.post(
                "/api/set-initial-password",
                json={
                    "token": setup_token,
                    "temporary_password": result["demo_temporary_password"],
                    "new_password": "DemoDeanPermanent1!",
                },
            )

        assert setup.status_code == 200, setup.text
        dean = db_session.query(models.User).filter_by(
            email="demo.dean@example.edu"
        ).one()
        assert verify_password("DemoDeanPermanent1!", dean.password)
        assert dean.password_setup_token_hash is None
    finally:
        app.dependency_overrides.clear()
        app.dependency_overrides.update(previous)


def test_campus_admin_can_resend_setup_for_dean_only_within_campus(db_session, monkeypatch):
    campus = models.Campus(name="Dean Resend Campus", code="DEAN-RESEND")
    other_campus = models.Campus(name="Other Dean Resend Campus", code="OTHER-DEAN-RESEND")
    db_session.add_all([campus, other_campus])
    db_session.flush()
    department = models.Department(name="Dean Resend Department", campus_id=campus.id)
    failed_department = models.Department(name="Failed Dean Department", campus_id=campus.id)
    other_department = models.Department(name="Foreign Dean Department", campus_id=other_campus.id)
    db_session.add_all([department, failed_department, other_department])
    db_session.flush()
    admin = models.User(
        email="dean-resend-admin@example.edu",
        password=hash_password("CampusAdminPassword1!"),
        role="campus_admin",
        campus_id=campus.id,
        archived=False,
    )
    dean = models.User(
        email="dean-resend@example.edu",
        password=hash_password("OldDeanPassword1!"),
        role="department_admin",
        campus_id=campus.id,
        admin_department_id=department.id,
        department=department.name,
        name="Dean Resend",
        archived=False,
    )
    foreign_dean = models.User(
        email="foreign-dean@example.edu",
        password=hash_password("OldDeanPassword1!"),
        role="department_admin",
        campus_id=other_campus.id,
        admin_department_id=other_department.id,
        department=other_department.name,
        name="Foreign Dean",
        archived=False,
    )
    failed_dean = models.User(
        email="failed-dean@example.edu",
        password=hash_password("OldDeanPassword1!"),
        role="department_admin",
        campus_id=campus.id,
        admin_department_id=failed_department.id,
        department=failed_department.name,
        name="Failed Dean",
        archived=False,
    )
    db_session.add_all([admin, dean, foreign_dean, failed_dean])
    db_session.flush()
    department.dean_id = dean.id
    failed_department.dean_id = failed_dean.id
    other_department.dean_id = foreign_dean.id
    db_session.commit()

    sent_credentials = []
    monkeypatch.setattr(main, "require_email_delivery_configured", lambda: None)
    monkeypatch.setattr(main, "send_department_faculty_access_email", lambda *args: sent_credentials.append(args) or True)
    monkeypatch.setattr(main, "log_activity", lambda *args, **kwargs: None)
    previous = with_test_dependencies(db_session, admin)
    try:
        with TestClient(app) as client:
            sent = client.post(
                f"/api/campus-admin/users/{dean.id}/credential-email",
                json={"action": "reset_password"},
            )
            denied = client.post(
                f"/api/campus-admin/users/{foreign_dean.id}/credential-email",
                json={"action": "reset_password"},
            )

            previous_password = failed_dean.password
            previous_token = failed_dean.password_setup_token_hash
            monkeypatch.setattr(main, "send_department_faculty_access_email", lambda *args: False)
            failed = client.post(
                f"/api/campus-admin/users/{failed_dean.id}/credential-email",
                json={"action": "reset_password"},
            )

        assert sent.status_code == 200, sent.text
        assert sent.json()["email_status"] == "sent"
        assert sent_credentials[-1][0] == dean.email
        assert sent_credentials[-1][3] == department.name
        assert denied.status_code == 404
        assert failed.status_code == 503
        db_session.refresh(failed_dean)
        assert failed_dean.password == previous_password
        assert failed_dean.password_setup_token_hash == previous_token
    finally:
        app.dependency_overrides.clear()
        app.dependency_overrides.update(previous)


def test_campus_admin_can_reset_password_only_for_faculty_in_assigned_campus(db_session, monkeypatch):
    campus = models.Campus(name="Reset Campus", code="RESET-CAMPUS")
    other_campus = models.Campus(name="Other Reset Campus", code="OTHER-RESET")
    db_session.add_all([campus, other_campus])
    db_session.flush()
    admin = models.User(
        email="reset-campus-admin@example.edu",
        password=hash_password("CampusAdminPassword1!"),
        role="campus_admin",
        campus_id=campus.id,
        archived=False,
    )
    faculty = models.User(
        email="reset-faculty@example.edu",
        password=hash_password("OriginalFacultyPassword1!"),
        role="faculty",
        name="Reset Faculty",
        department="Reset Department",
        campus_id=campus.id,
        archived=False,
    )
    foreign_faculty = models.User(
        email="foreign-faculty@example.edu",
        password=hash_password("OriginalFacultyPassword1!"),
        role="faculty",
        name="Foreign Faculty",
        department="Foreign Department",
        campus_id=other_campus.id,
        archived=False,
    )
    db_session.add_all([admin, faculty, foreign_faculty])
    db_session.commit()

    sent_credentials = []
    monkeypatch.setattr(main, "require_email_delivery_configured", lambda: None)
    monkeypatch.setattr(
        main,
        "send_department_faculty_access_email",
        lambda *args: sent_credentials.append(args) or True,
    )
    monkeypatch.setattr(main, "log_activity", lambda *args, **kwargs: None)
    previous = with_test_dependencies(db_session, admin)
    try:
        with TestClient(app) as client:
            reset = client.post(
                f"/api/campus-admin/users/{faculty.id}/credential-email",
                json={"action": "reset_password"},
            )
            denied = client.post(
                f"/api/campus-admin/users/{foreign_faculty.id}/credential-email",
                json={"action": "reset_password"},
            )

        assert reset.status_code == 200, reset.text
        assert reset.json()["email_status"] == "sent"
        assert denied.status_code == 404
        db_session.refresh(faculty)
        assert verify_password(sent_credentials[0][1], faculty.password)
        assert faculty.password_setup_token_hash
        assert sent_credentials[0][0] == faculty.email
        assert sent_credentials[0][3] == faculty.department
        assert len(sent_credentials) == 1
    finally:
        app.dependency_overrides.clear()
        app.dependency_overrides.update(previous)


def test_department_admin_creates_program_chair_account_and_emails_credentials(db_session, monkeypatch):
    campus = models.Campus(name="Chair Campus", code="CHAIR")
    db_session.add(campus)
    db_session.flush()
    department = models.Department(name="Chair Department", campus_id=campus.id)
    other_department = models.Department(name="Other Chair Department", campus_id=campus.id)
    db_session.add_all([department, other_department])
    db_session.flush()
    program = models.Program(name="Chair Program", department_id=department.id)
    duplicate_test_program = models.Program(
        name="Duplicate Test Program",
        department_id=department.id,
    )
    other_program = models.Program(name="Other Program", department_id=other_department.id)
    admin = models.User(
        email="chair-admin@example.com",
        role="department_admin",
        campus_id=campus.id,
        admin_department_id=department.id,
        archived=False,
    )
    db_session.add_all([program, duplicate_test_program, other_program, admin])
    db_session.commit()

    delivered = []
    monkeypatch.setattr("main.require_email_delivery_configured", lambda: None)
    monkeypatch.setattr("main.generate_temporary_password", lambda: "ChairPassword1!")
    monkeypatch.setattr("main.log_activity", lambda *args, **kwargs: None)
    monkeypatch.setattr(
        "main.send_program_chair_credentials_email",
        lambda *args: delivered.append(args),
    )
    previous = with_test_dependencies(db_session, admin)
    try:
        with TestClient(app) as client:
            response = client.post(
                f"/api/programs/{program.id}/chair-account",
                json={"full_name": "Alex Chair", "email": "Alex.Chair@example.edu"},
            )
            assert response.status_code == 201, response.text
            result = response.json()
            created_user = db_session.query(models.User).filter_by(
                email="alex.chair@example.edu"
            ).one()
            db_session.refresh(program)
            assert created_user.role == "faculty"
            assert created_user.program_id == program.id
            assert created_user.campus_id == campus.id
            assert created_user.department == department.name
            assert verify_password("ChairPassword1!", created_user.password)
            assert program.chair_id == created_user.id
            assert program.chair_name == "Alex Chair"
            assert result["email_status"] == "queued"
            assert len(delivered) == 1
            assert delivered[0] == (
                "alex.chair@example.edu",
                "Alex Chair",
                "ChairPassword1!",
                program.name,
                department.name,
            )

            already_assigned = client.post(
                f"/api/programs/{program.id}/chair-account",
                json={"full_name": "Second Chair", "email": "second.chair@example.edu"},
            )
            assert already_assigned.status_code == 409

            duplicate_email = client.post(
                f"/api/programs/{duplicate_test_program.id}/chair-account",
                json={"full_name": "Duplicate Chair", "email": "ALEX.CHAIR@example.edu"},
            )
            assert duplicate_email.status_code == 409

            denied_program = client.post(
                f"/api/programs/{other_program.id}/chair-account",
                json={"full_name": "Other Chair", "email": "other.chair@example.edu"},
            )
            assert denied_program.status_code == 403
    finally:
        app.dependency_overrides.clear()
        app.dependency_overrides.update(previous)


def test_campus_admin_can_review_department_academic_change_requests(db_session, monkeypatch):
    campus = models.Campus(name="Review Campus", code="REV")
    db_session.add(campus)
    db_session.flush()
    department = models.Department(name="Review Department", campus_id=campus.id)
    db_session.add(department)
    db_session.flush()
    campus_admin = models.User(
        email="campus-admin@example.com",
        role="campus_admin",
        campus_id=campus.id,
        archived=False,
    )
    department_admin = models.User(
        email="department-admin@example.com",
        role="department_admin",
        campus_id=campus.id,
        admin_department_id=department.id,
        archived=False,
    )
    db_session.add_all([campus_admin, department_admin])
    db_session.flush()
    request_entry = models.DepartmentAcademicChangeRequest(
        department_id=department.id,
        submitted_by=department_admin.id,
        title="Review curriculum mapping",
        details="Please check the program subject mapping.",
        status="pending",
    )
    db_session.add(request_entry)
    db_session.commit()
    monkeypatch.setattr("main.log_activity", lambda *args, **kwargs: None)

    previous = with_test_dependencies(db_session, campus_admin)
    try:
        with TestClient(app) as client:
            listed = client.get("/api/department-academic-change-requests")
            assert listed.status_code == 200
            assert [row["id"] for row in listed.json()] == [request_entry.id]
            reviewed = client.put(
                f"/api/department-academic-change-requests/{request_entry.id}",
                json={"status": "needs_info", "response": "Please include the affected semester."},
            )
            assert reviewed.status_code == 200
            db_session.refresh(request_entry)
            assert request_entry.status == "needs_info"
            assert request_entry.response == "Please include the affected semester."

            app.dependency_overrides[get_current_user] = lambda: department_admin
            response = client.put(
                f"/api/department-academic-change-requests/{request_entry.id}/respond",
                json={"details": "The affected semesters are the first and second."},
            )
            assert response.status_code == 200
            db_session.refresh(request_entry)
            assert request_entry.status == "pending"
            assert request_entry.details == "The affected semesters are the first and second."

            app.dependency_overrides[get_current_user] = lambda: campus_admin
            approved = client.put(
                f"/api/department-academic-change-requests/{request_entry.id}",
                json={"status": "approved"},
            )
            assert approved.status_code == 200
            db_session.refresh(request_entry)
            assert request_entry.status == "approved"
    finally:
        app.dependency_overrides.clear()
        app.dependency_overrides.update(previous)
