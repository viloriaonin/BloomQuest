import openpyxl
wb = openpyxl.load_workbook('templates/tos_template.xlsx', data_only=False)
print(wb.sheetnames)
ws = wb.active
print('active', ws.title)
for r in range(1, 40):
    row = [ws.cell(row=r, column=c).value for c in range(1, 26)]
    if any(x is not None for x in row):
        print(r, row)
